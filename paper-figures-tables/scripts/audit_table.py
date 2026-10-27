"""审已有表：只查**表格层**，不查整页版式。

用途是「学生把表做好了，看看能不能交」。查的是表题、三条横线、表头下线、
字号区间、前导零、小数位、统计符号斜体——全部是图表本身的事。

**不查**页边距、行距、页眉页脚、章节编号、图表目录。那些属于整篇论文的排版，
超出本技能「图表的制作与质量检查」的边界；硬查会带来大量「该刊/该校本来就这样」
的误报——同一篇论文按 APA 审，风格差异能报出几十条。

用法：

    python scripts/audit_table.py paper.docx
    python scripts/audit_table.py paper.docx --profile zh-journal --json report.json
    python scripts/audit_table.py paper.docx --table 2      # 只审第 2 张表
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import profiles  # noqa: E402
import table_checks as tc  # noqa: E402
from qc import E, I, W, Finding, print_findings, write_report  # noqa: E402


def iter_blocks(doc):
    """按文档顺序产出段落与表格。python-docx 的 .paragraphs / .tables 会丢掉相对顺序。"""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            yield Table(child, doc)


def audit(path: Path, profile: dict, only: int | None = None) -> list[Finding]:
    from docx import Document

    # 表格层的判据（边框、合并单元格、字号、斜体）都写在 docx 的 XML 里，
    # PDF 里没有这些信息——拿 PDF 来审只能读文本层，读不出线型与合并。
    # 与其崩在 python-docx 的 PackageNotFoundError 上，不如说清楚该给什么。
    if path.suffix.lower() not in (".docx", ".doc"):
        bad = [Finding("PFT-TBL-000", E,
                       f"{path.name} 不是 Word 文档：审表只能读 .docx",
                       limit="PDF 里没有边框、合并单元格、字号这些信息；"
                             "要审 PDF 里的表，先用 make_table 复刻成 docx，或向作者要原始 docx")]
        print_findings(bad)
        return bad

    try:
        doc = Document(str(path))
    except Exception as e:                      # noqa: BLE001
        bad = [Finding("PFT-TBL-000", E,
                       f"{path.name} 打不开（{type(e).__name__}）：可能不是有效的 docx",
                       limit="只支持 python-docx 能解析的 .docx")]
        print_findings(bad)
        return bad
    # 本技能生成的表会在文档属性里声明「哪些列显式保留前导零」。
    # 读它，才能让「审自己生成的表」与「审别人的表」用同一套判据而不误报。
    keep_lz: set[str] = set()
    for part in (doc.core_properties.comments or "").split(";"):
        if part.startswith("pft-lz:"):
            keep_lz = {s.strip() for s in part.split(":", 1)[1].split(",") if s.strip()}
    blocks = list(iter_blocks(doc))
    idx = [i for i, b in enumerate(blocks) if hasattr(b, "rows") and hasattr(b, "columns")]
    if not idx:
        return [Finding("PFT-TBL-000", E, f"{path.name} 里没有表格")]

    findings: list[Finding] = []
    for k, i in enumerate(idx, 1):
        if only is not None and k != only:
            continue
        t = blocks[i]
        prevs, after = tc.caption_context(blocks, i)
        h_end = tc.header_end(t)
        is_panel = "pft-panels:1" in (doc.core_properties.comments or "")
        got = (tc.caption_findings(prevs, after, profile)
               + tc.border_findings(t)
               + tc.width_findings(t)
               + tc.page_fit_findings(t, doc)
               + tc.negative_zero_findings(t)
               + tc.placeholder_findings(t)
               + tc.word_break_findings(t, profile))
        if not is_panel:
            # 多面板表没有唯一表头行，这几项按「表头 + 表体」划分，套上去只会误报
            got += (tc.header_rule_findings(t, h_end)
                    + tc.font_size_findings(t, profile)
                    + tc.stat_symbol_findings(t)
                    + tc.leading_zero_findings(t, h_end, profile, keep_lz)
                    + tc.decimals_findings(t, h_end))
        got += tc.empty_run_findings(t)
        label = f"第 {k} 张表（{len(t.rows)} 行 × {len(t.columns)} 列）"
        print(f"\n{label}")
        if not got:
            print("  ✓ 表格层检查通过")
        else:
            print_findings(got)
        findings += got

    print("\n未检查的范围（需人工）：页边距与整页版式、行距、页眉页脚、"
          "表内每个缩写是否在表注中定义、表里的数字与原始输出是否一致、"
          "表是否在正文中被引用。")
    return findings


def main() -> int:
    ap = argparse.ArgumentParser(description="审已有表（表格层）")
    ap.add_argument("files", nargs="+")
    ap.add_argument("--profile", help="按哪个 profile 判（默认 apa7）")
    ap.add_argument("--table", type=int, help="只审第 N 张表（从 1 起）")
    ap.add_argument("--json", help="把结论写成 JSON 报告")
    args = ap.parse_args()

    profile = profiles.load(args.profile or "apa7")
    all_findings: list[Finding] = []
    for name in args.files:
        p = Path(name)
        print(f"\n=== {p.name}（profile: {profile['profile_id']}）")
        if not p.exists():
            all_findings.append(Finding("PFT-GEN-404", E, f"{p.name} 不存在"))
            continue
        all_findings += audit(p, profile, args.table)

    n_err = sum(1 for f in all_findings if f.severity == E)
    n_warn = sum(1 for f in all_findings if f.severity == W)
    print(f"\n合计 {n_err} 个必须修的问题，{n_warn} 个提示")
    print("本报告不构成期刊合规认证：只覆盖规则目录里列明的自动检查项。")
    if args.json:
        out = write_report(args.json, all_findings, targets=args.files,
                           tool="paper-figures-tables/audit_table")
        print(f"报告：{out}")
    return 1 if n_err else 0


if __name__ == "__main__":
    sys.exit(main())
