"""技能自检：把每个生成器跑一遍，检查产出是否符合规则。

用法：
    python scripts/selftest.py            # 全部跑一遍
    python scripts/selftest.py -v         # 同时打印每个用例的细节

它做的事：
  1. 在临时目录里造规格，跑五个生成器（表格 / 示意图 / 交互图 / 估计图 / 多面板）
  2. 用 check_figure.py 与 check_connections.py 的判据复查产出
  3. 故意写错几处，确认会**报错而不是静默出错**
  4. 同一输入跑两次，比对 md5，确认构建是确定的

为什么要有这个文件：改一个生成器很容易把另一个弄坏，而「图看起来还行」
骗得过眼睛。这个脚本是唯一能一次验证全部规则的地方。
"""
from __future__ import annotations

import hashlib
import json
import math
import random
import re
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))

import audit_figure  # noqa: E402
import audit_table  # noqa: E402
import check_figure  # noqa: E402
import make_diagram  # noqa: E402
import make_estimate_plot as mep  # noqa: E402
import make_interaction  # noqa: E402
import make_line_plot as line_plot  # noqa: E402
import make_panel  # noqa: E402
import make_table  # noqa: E402
import preflight as pf  # noqa: E402
import profiles  # noqa: E402
from qc import Finding, write_report  # noqa: E402

VERBOSE = "-v" in sys.argv
RESULTS: list[tuple[str, bool, str]] = []


def record(name: str, ok: bool, note: str = "") -> None:
    RESULTS.append((name, ok, note))
    mark = "✓" if ok else "✗"
    print(f"  {mark} {name}" + (f"  —— {note}" if note and (not ok or VERBOSE) else ""))


def md5(p: Path) -> str:
    return hashlib.md5(p.read_bytes()).hexdigest()


def _pdf_width(path: Path) -> float:
    """读 PDF 首页宽度（pt）。没有 PyMuPDF 时返回 -1，让断言如实失败。"""
    try:
        import pymupdf
        with pymupdf.open(str(path)) as d:
            return float(d[0].rect.width)
    except Exception:                                   # noqa: BLE001
        return -1.0


def expect_exit(fn, needle: str) -> tuple[bool, str]:
    """调用 fn，期望它抛 SystemExit 且消息里含 needle。"""
    try:
        fn()
    except SystemExit as e:
        msg = str(e)
        return (needle in msg), msg.splitlines()[-1] if msg else ""
    except Exception as e:                      # noqa: BLE001
        return False, f"抛的是 {type(e).__name__} 而不是 SystemExit：{e}"
    return False, "没有报错"


# ---------------------------------------------------------------- 夹具

TABLE_APA = {
    "number": "1",
    "title": "Descriptive Statistics and Correlations",
    "format": "apa7",
    "columns": [
        {"key": "var", "label": "Variable"},
        {"key": "m", "label": "M", "decimals": 2, "leading_zero": True},
        {"key": "sd", "label": "SD", "decimals": 2, "leading_zero": True},
        {"key": "r", "label": "1", "align": "decimal"},
        {"key": "r2", "label": "2", "align": "decimal"},
    ],
    "rows": [
        {"var": "1. Trust", "m": "4.2", "sd": "1.1", "r": ".85", "r2": ""},
        {"var": "2. Commitment", "m": "3.97", "sd": "1.14", "r": ".53**", "r2": ".83"},
    ],
    "notes": {"general": "N = 344.", "probability": "**p < .01."},
}

TABLE_ZH = {
    "number": "2",
    "title": "变量描述统计与相关系数",
    "format": "zh",
    "columns": [
        {"key": "var", "label": "变量"},
        {"key": "m", "label": "均值", "decimals": 2, "leading_zero": True},
        {"key": "a", "label": "α"},
        {"key": "r1", "label": "1"},
        {"key": "r2", "label": "2"},
    ],
    "spanners": [{"label": "信度", "span": ["a"]},
                 {"label": "相关系数", "span": ["r1", "r2"]}],
    "rows": [
        {"var": "1. 品牌信任", "m": "4.21", "a": "0.89", "r1": "0.85", "r2": ""},
        {"var": "2. 品牌承诺", "m": "3.97", "a": "0.87", "r1": "0.53**", "r2": "0.83"},
    ],
    "notes": {"general": "N = 344。"},
}

DIAG_CONCEPT = {
    "type": "conceptual",
    "lang": "en",
    "title": "Figure 1",
    "nodes": [
        {"id": "a", "label": "Shared values"},
        {"id": "b", "label": "Communication"},
        {"id": "t", "label": "Trust", "level": 1},
        {"id": "c", "label": "Commitment", "level": 2},
        {"id": "o", "label": "Cooperation", "level": 3},
    ],
    "edges": [
        {"from": "a", "to": "t", "label": "H1"},
        {"from": "b", "to": "t", "label": "H2"},
        {"from": "t", "to": "c", "label": "H3"},
        {"from": "c", "to": "o", "label": "H4"},
    ],
    "moderators": [{"id": "u", "label": "Uncertainty", "on": ["t", "c"],
                    "label_on_edge": "H5"}],
}

DIAG_MEAS = {
    "type": "measurement",
    "lang": "zh",
    "row_gap": 40,
    "nodes": [
        {"id": "tr", "label": "品牌信任", "shape": "ellipse", "level": 0},
        {"id": "t1", "label": "值得信赖", "level": 1},
        {"id": "t2", "label": "言行一致", "level": 1},
        {"id": "cm", "label": "品牌承诺", "shape": "ellipse", "level": 0},
        {"id": "c1", "label": "情感依恋", "level": 1},
    ],
    "edges": [
        {"from": "tr", "to": "t1", "coef": ".82"},
        {"from": "tr", "to": "t2", "coef": ".76"},
        {"from": "cm", "to": "c1", "coef": ".88"},
        {"from": "tr", "to": "cm", "bidirectional": True, "coef": ".53***"},
    ],
}


def make_interaction_csv(path: Path, n: int = 240, b_int: float = 0.40) -> None:
    rng = random.Random(20261025)
    rows = ["x,m,y"]
    for _ in range(n):
        x = rng.gauss(4.0, 1.0)
        m = rng.gauss(3.0, 1.0)
        y = 2.0 + 0.5 * x + 0.3 * m + b_int * x * m + rng.gauss(0, 0.8)
        rows.append(f"{x:.4f},{m:.4f},{y:.4f}")
    path.write_text("\n".join(rows) + "\n", encoding="utf-8", newline="\n")


def _make_bad_table(path: Path) -> None:
    """造一张「学生常见写法」的表：无表题、全框线、表头符号不斜体、字号混杂、小数位不一。"""
    from docx import Document
    from docx.shared import Pt
    d = Document()
    d.add_paragraph("这是一段正文，表格前面没有表题。")
    t = d.add_table(rows=4, cols=3)
    for c, s in zip(t.rows[0].cells, ["变量", "M", "r"]):
        r = c.paragraphs[0].add_run(s)
        r.font.size = Pt(10.5)
    for i, (a, b, cc) in enumerate([("信任", "4.21", "0.53"),
                                    ("承诺", "3.97", ".5"),
                                    ("忠诚", "4.02", ".612")], 1):
        for j, v in enumerate((a, b, cc)):
            r = t.rows[i].cells[j].paragraphs[0].add_run(v)
            r.font.size = Pt(10.5 if j else 9)
    d.add_paragraph("注：数据为示例。")
    d.save(str(path))


EST_DOT = {
    "lang": "en",
    "width": 220,
    "xlabel": "Mean (95% CI)",
    "refline": 4,
    "refline_label": "Midpoint",
    "series": [
        {"label": "", "points": [
            {"label": "Brand A", "est": 4.21, "lo": 3.98, "hi": 4.44, "sig": "**"},
            {"label": "Brand B", "est": 3.80, "lo": 3.55, "hi": 4.05},
            {"label": "Brand C", "est": 4.63, "lo": 4.41, "hi": 4.85, "sig": "*"},
        ]},
    ],
}

EST_GROUPED = {
    "lang": "en",
    "xlabel": "Standardized coefficient (95% CI)",
    "refline": 0,
    "refline_label": "No effect",
    "series": [
        {"label": "Study 1", "points": [
            {"label": "Trust → Commitment", "est": 0.53, "lo": 0.44, "hi": 0.62},
            {"label": "Trust → Loyalty", "est": 0.09, "lo": -0.02, "hi": 0.20},
        ]},
        {"label": "Study 2", "points": [
            {"label": "Trust → Commitment", "est": 0.47, "lo": 0.38, "hi": 0.56},
            {"label": "Trust → Loyalty", "est": 0.14, "lo": 0.03, "hi": 0.25},
        ]},
    ],
}

LONG_TABLE = {
    "number": "4",
    "title": "Measurement Model: Factor Loadings",
    "format": "apa7",
    "longtable": True,
    "columns": [{"key": "item", "label": "Item"},
                {"key": "load", "label": "λ", "decimals": 3},
                {"key": "se", "label": "SE", "decimals": 3}],
    "rows": [{"item": f"Item {i}", "load": f"0.{70 + i % 20}", "se": "0.05"}
             for i in range(1, 26)],
    "notes": {"general": "N = 344.", "specific": ["Fixed parameter."]},
}

RICH_TABLE = {
    "number": "2",
    "title": "Hierarchical Regression Predicting Loyalty",
    "format": "apa7",
    "columns": [
        {"key": "var", "label": "Predictor"},
        {"key": "b", "label": "<i>B</i>", "decimals": 2, "leading_zero": True},
        {"key": "se", "label": "<i>SE</i>", "decimals": 2, "leading_zero": True},
        {"key": "beta", "label": "β", "decimals": 2},
        {"key": "p", "label": "<i>p</i>"},
    ],
    "rows": [
        {"var": "Step 1", "b": "", "se": "", "beta": "", "p": ""},
        {"var": "  Trust", "b": "0.52", "se": "0.07", "beta": "0.48", "p": "< .001"},
        {"var": "  Marked", "b": "<i>0.31</i>", "se": "0.06", "beta": "0.29", "p": "< .001"},
    ],
    "notes": {"general": "<i>N</i> = 344. <i>B</i> = unstandardized coefficient.",
              "specific": ["Reference group = low trust."],
              "probability": "*<i>p</i> < .05. **<i>p</i> < .01. ***<i>p</i> < .001."},
}


def all_texts(path: Path) -> list[tuple[float, float, str]]:
    import xml.etree.ElementTree as ET
    root = ET.parse(path).getroot()
    out = []
    for el in root.iter("{http://www.w3.org/2000/svg}text"):
        out.append((float(el.get("y") or 0), float(el.get("x") or 0), el.text or ""))
    return out


def horizontal_line_ys(path: Path) -> list[float]:
    import xml.etree.ElementTree as ET
    root = ET.parse(path).getroot()
    ys = []
    for el in root.iter("{http://www.w3.org/2000/svg}line"):
        if el.get("y1") == el.get("y2"):
            ys.append(float(el.get("y1")))
    return ys


# ---------------------------------------------------------------- 用例

def run(tmp: Path) -> None:
    # ---- 表格
    for name, spec, fmt_ok in (("表格 APA", TABLE_APA, True), ("表格 中文（分组表头）", TABLE_ZH, True)):
        j = tmp / f"{name}.json"
        j.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
        docx = tmp / f"{name}.docx"
        make_table.build_docx(spec, docx)
        errs, warns = check_figure.check_docx(docx)
        record(f"{name}：生成并自检", not errs, "；".join(errs) or f"{len(warns)} 个提示")

    # 小数位归一
    docx = tmp / "表格 APA.docx"
    try:
        from docx import Document
        t = Document(str(docx)).tables[0]
        vals = [t.cell(r, 1).text for r in (1, 2)]
        ok = vals == ["4.20", "3.97"]
        record("表格：decimals 把 4.2 补成 4.20", ok, " / ".join(vals))
    except ImportError:
        record("表格：decimals 归一", False, "缺 python-docx，跳过")

    # ---- 示意图
    geo = {}
    svg, warns = make_diagram.build(DIAG_CONCEPT, geo)
    svg.save(str(tmp / "概念.svg"))
    (tmp / "概念.geo.json").write_text(json.dumps(geo, ensure_ascii=False), encoding="utf-8")
    errs, _ = check_figure.check_svg(tmp / "概念.svg")
    record("概念框架图：生成并自检", not errs, "；".join(errs))
    record("概念框架图：无校验警告", not warns, "；".join(warns))
    bad = connection_problems(geo)
    record("概念框架图：净空/穿框/编号归属", not bad, "；".join(bad[:3]))

    geo2 = {}
    svg2, warns2 = make_diagram.build(DIAG_MEAS, geo2)
    svg2.save(str(tmp / "测量.svg"))
    record("测量模型：无校验警告", not warns2, "；".join(warns2))
    # 双向箭头的实现是「两条坐标互为反向的线」——直接查这个特征，
    # 比数线条数量可靠：载荷线也会带 marker-end。
    text = (tmp / "测量.svg").read_text(encoding="utf-8")
    coords = set()
    for m in re.finditer(r'<line x1="([-\d.]+)" y1="([-\d.]+)" '
                         r'x2="([-\d.]+)" y2="([-\d.]+)"[^>]*marker-end', text):
        coords.add(tuple(m.groups()))
    pairs = [c for c in coords if (c[2], c[3], c[0], c[1]) in coords]
    record("测量模型：潜变量之间画了双向箭头", bool(pairs),
           f"{len(pairs)//2} 条双向线（{len(coords)} 条带箭头的线）")
    bad2 = connection_problems(geo2)
    record("测量模型：净空/穿框/编号归属", not bad2, "；".join(bad2[:3]))

    # ---- 交互图
    csv = tmp / "inter.csv"
    make_interaction_csv(csv)
    df = make_interaction._read_csv(str(csv))
    fit = make_interaction.fit_interaction(df, "y", "x", "m")
    lo, hi = fit["x_mean"] - 2 * fit["x_sd"], fit["x_mean"] + 2 * fit["x_sd"]
    grid = [lo + (hi - lo) * i / 59 for i in range(60)]
    series = []
    for lv, lab in ((fit["m_mean"] - fit["m_sd"], "Low"),
                    (fit["m_mean"] + fit["m_sd"], "High")):
        yhat, se = make_interaction.predict(fit, grid, lv)
        series.append({"label": lab, "x": grid, "y": yhat,
                       "lo": [a - 1.96 * s for a, s in zip(yhat, se)],
                       "hi": [a + 1.96 * s for a, s in zip(yhat, se)]})
    svg3 = make_interaction.build({"series": series},
                                  {"lang": "en", "xlabel": "X", "ylabel": "Y"})
    svg3.save(str(tmp / "交互.svg"))
    errs3, _ = check_figure.check_svg(tmp / "交互.svg")
    record("交互效应图：生成并自检", not errs3, "；".join(errs3))
    b = fit["beta"][3]
    record("交互效应图：交互系数接近真值 0.40", abs(b - 0.40) < 0.08, f"b = {b:.3f}")

    # ---- 轴刻度必须在轴的外侧（回归：刻度文字曾落在绘图区里）
    ys = horizontal_line_ys(tmp / "交互.svg")
    axis_y = max(ys) if ys else 0.0
    below = [t for t in all_texts(tmp / "交互.svg") if t[0] > axis_y + 2]
    record("交互效应图：x 轴刻度文字在轴的下方", bool(below),
           f"轴线 y={axis_y:.1f}，轴下方 {len(below)} 处文字")

    # ---- 估计图（组间均值 / 系数 + 误差棒）
    svg4, w4 = mep.build(EST_DOT)
    svg4.save(str(tmp / "估计点图.svg"))
    errs4, _ = check_figure.check_svg(tmp / "估计点图.svg")
    record("估计图：单系列生成并自检", not errs4, "；".join(errs4))
    record("估计图：无警告", not w4, "；".join(w4))

    svg5, _ = mep.build(EST_GROUPED)
    svg5.save(str(tmp / "估计分组.svg"))
    errs5, _ = check_figure.check_svg(tmp / "估计分组.svg")
    record("估计图：两组并排生成并自检", not errs5, "；".join(errs5))

    got = [mep.t_crit(df) for df in (1, 10, 30, 1000)]
    want = [12.706, 2.228, 2.042, 1.962]
    record("估计图：t 分位数（df=1/10/30/1000）",
           all(abs(a - b) < 5e-3 for a, b in zip(got, want)),
           " / ".join(f"{g:.3f}" for g in got))

    bar_spec = json.loads(json.dumps(EST_DOT))
    bar_spec["style"] = "bar"
    bar_spec.pop("refline")
    svg6, _ = mep.build(bar_spec)
    svg6.save(str(tmp / "估计柱状.svg"))
    errs6, _ = check_figure.check_svg(tmp / "估计柱状.svg")
    record("估计图：柱状图无 refline 时基线纳入定义域", not errs6, "；".join(errs6))

    csv2 = tmp / "est.csv"
    csv2.write_text("cond,y\nA,1\nA,2\nA,3\nB,4\nB,5\nB,6\n", encoding="utf-8")
    df2 = mep._read_csv(str(csv2))
    ser, _ = mep.summarize(df2, "y", "cond")
    p0 = ser[0]["points"][0]
    t2 = mep.t_crit(2)
    want_lo, want_hi = 2 - t2 / math.sqrt(3), 2 + t2 / math.sqrt(3)
    ok = (abs(p0["est"] - 2) < 1e-9 and abs(p0["lo"] - want_lo) < 1e-9
          and abs(p0["hi"] - want_hi) < 1e-9)
    record("估计图：从 CSV 汇总的均值与 t 区间", ok,
           f"M={p0['est']:.3f} [{p0['lo']:.3f}, {p0['hi']:.3f}]，"
           f"手算 [{want_lo:.3f}, {want_hi:.3f}]")

    bad_est = {
        "区间不含估计值": ({**EST_DOT, "series": [{"points": [
            {"label": "A", "est": 1.0, "lo": 2.0, "hi": 3.0}]}]}, "区间不包含点估计"),
        "缺 lo": ({**EST_DOT, "series": [{"points": [
            {"label": "A", "est": 1.0, "hi": 3.0}]}]}, "缺数值 `lo`"),
        "把 est 拼成 mean": ({**EST_DOT, "series": [{"points": [
            {"label": "A", "mean": 1.0, "lo": 0.0, "hi": 2.0}]}]}, "未知键"),
        "各组点数不一致": ({**EST_DOT, "series": [
            {"points": [{"label": "A", "est": 1, "lo": 0, "hi": 2}]},
            {"points": [{"label": "A", "est": 1, "lo": 0, "hi": 2},
                        {"label": "B", "est": 1, "lo": 0, "hi": 2}]}]}, "一一对应"),
    }
    for label, (spec, needle) in bad_est.items():
        ok, msg = expect_exit(lambda s=spec: mep.build(s), needle)
        record(f"负向：估计图{label} 应报错", ok, msg)

    # ---- 多面板图
    for nm, spec in (("pa", EST_DOT), ("pb", {**EST_DOT, "style": "bar"})):
        s, _ = mep.build(spec)
        s.save(str(tmp / f"{nm}.svg"))
    panels = [make_panel.read_panel(tmp / "pa.svg"), make_panel.read_panel(tmp / "pb.svg")]
    svg7, w7, info7 = make_panel.build(panels, {"labels": ["(a)", "(b)"],
                                               "note": "Error bars are 95% CIs."})
    svg7.save(str(tmp / "面板.svg"))
    errs7, _ = check_figure.check_svg(tmp / "面板.svg")
    record("多面板：合成两幅并自检", not errs7, "；".join(errs7))
    record("多面板：宽度不超过 6.5 in", info7["width_in"] <= 6.51,
           f"{info7['width_in']:.2f} in，缩放 {info7['scale']:.3f}")
    txt7 = (tmp / "面板.svg").read_text(encoding="utf-8")
    record("多面板：(a)(b) 标注都在", "(a)" in txt7 and "(b)" in txt7)

    # 合成的图必须把面板内容写进 ops，否则 matplotlib 重画 PDF 时整页空白
    # （SVG 却完全正常——这是只有打开导出的 PDF 才看得见的缺陷）
    panel_texts = [op["s"] for p in panels for op in p["ops"] if op["op"] == "text"]
    comp_texts = [op["s"] for op in svg7.ops if op["op"] == "text"]
    missing = [t for t in panel_texts if t not in comp_texts]
    record("多面板：面板的文字都进了合成图的 ops（否则导出 PDF 是空白）",
           not missing and len(svg7.ops) > len(panels[0]["ops"]),
           f"缺 {len(missing)} 条，ops {len(svg7.ops)} 条")

    big = [make_panel.read_panel(tmp / "概念.svg"), make_panel.read_panel(tmp / "测量.svg")]
    svg8, w8, info8 = make_panel.build(big, {"labels": ["(a)", "(b)"]})
    svg8.save(str(tmp / "面板大.svg"))
    record("多面板：超宽时等比缩小并警告",
           info8["scale"] < 1.0 and any("最小字号" in x for x in w8),
           f"scale={info8['scale']:.3f}；" + "；".join(w8))
    txt8 = (tmp / "面板大.svg").read_text(encoding="utf-8")
    record("多面板：marker 定义去重", txt8.count('id="ar"') == 1,
           f'id="ar" 出现 {txt8.count(chr(34) + "ar" + chr(34))} 次')

    # ---- 长表跨页
    make_table.build_docx(LONG_TABLE, tmp / "长表.docx")
    try:
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(str(tmp / "长表.docx"))
        rows = doc.tables[0].rows
        trPr = rows[0]._tr.find(qn("w:trPr"))
        record("表格：表头行设了 tblHeader（跨页重复）",
               trPr is not None and trPr.find(qn("w:tblHeader")) is not None)
        trPr2 = rows[1]._tr.find(qn("w:trPr"))
        record("表格：数据行设了 cantSplit（不跨页断行）",
               trPr2 is not None and trPr2.find(qn("w:cantSplit")) is not None)
        errs_l, _ = check_figure.check_docx(tmp / "长表.docx")
        record("表格：25 行长表自检通过", not errs_l, "；".join(errs_l))
    except ImportError:
        record("表格：长表结构", False, "缺 python-docx，跳过")
    make_table.build_tex(LONG_TABLE, tmp / "长表.tex")
    tex = (tmp / "长表.tex").read_text(encoding="utf-8")
    record("表格：--tex 长表用 longtable",
           "\\begin{longtable}" in tex and "\\endhead" in tex)

    # ---- 检查器本身：文字越界必须报错
    bad_svg = tmp / "越界.svg"
    bad_svg.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="100pt" height="50pt" '
        'viewBox="0 0 100 50"><text x="90" y="25" font-size="10">'
        'This label is far too long</text></svg>', encoding="utf-8")
    errs_b, _ = check_figure.check_svg(bad_svg)
    record("负向：文字超出画布应报错",
           any("超出画布" in e for e in errs_b), "；".join(errs_b))

    # ---- 富文本单元格：APA 的统计符号必须斜体
    make_table.build_docx(RICH_TABLE, tmp / "富文本.docx")
    try:
        from docx import Document
        d = Document(str(tmp / "富文本.docx"))
        hdr = d.tables[0].cell(0, 1).paragraphs[0].runs
        record("表格：<i> 标记排成斜体（表头）",
               len(hdr) == 1 and hdr[0].text == "B" and bool(hdr[0].italic),
               repr([(r.text, r.italic) for r in hdr]))
        note_runs = d.paragraphs[-1].runs
        record("表格：表注里的 <i> 也生效",
               any(r.text == "N" and r.italic for r in note_runs),
               repr([(r.text[:8], r.italic) for r in note_runs][:5]))
        marked = d.tables[0].cell(3, 1).paragraphs[0].runs
        record("表格：带标记的单元格不被数字格式化改坏",
               len(marked) == 1 and marked[0].text == "0.31" and bool(marked[0].italic),
               repr([(r.text, r.italic) for r in marked]))
        errs_r, _ = check_figure.check_docx(tmp / "富文本.docx")
        record("表格：富文本表自检通过", not errs_r, "；".join(errs_r))
    except ImportError:
        record("表格：富文本", False, "缺 python-docx，跳过")

    # ---- 图题不得压在图元上（自检夹具曾经真的重叠过）
    geo_t = {}
    make_diagram.build(DIAG_CONCEPT, geo_t)
    bad_t = connection_problems(geo_t)
    record("示意图：图题不压节点/调节变量",
           not any("图题压住" in b for b in bad_t), "；".join(bad_t[:2]))
    geo_bad = json.loads(json.dumps(geo_t))
    first = next(iter(geo_bad["nodes"]))
    geo_bad["title"]["box"] = geo_bad["nodes"][first]["box"]
    hits = connection_problems(geo_bad)
    record("负向：图题压住节点应报错",
           any("图题压住" in h for h in hits), "；".join(hits[:2]))

    # ---- matplotlib 导出后端（可选依赖）
    import importlib.util
    if importlib.util.find_spec("matplotlib"):
        files, notes = svg.export(tmp / "导出", formats=("pdf", "png"), dpi=150)
        ok = (len(files) == 2 and all(f.stat().st_size > 800 for f in files))
        record("导出：matplotlib 后端出 PDF 与 PNG", ok,
               f"{len(files)} 个文件；" + "；".join(notes))
        pdf = next((f for f in files if f.suffix == ".pdf"), None)
        if pdf:
            record("导出：PDF 页面尺寸与 SVG 一致（pt）",
                   abs(_pdf_width(pdf) - svg.w) < 0.5,
                   f"PDF {_pdf_width(pdf):.1f} pt vs SVG {svg.w:.1f} pt")
    else:
        record("导出：matplotlib 后端", True, "未安装 matplotlib，跳过（SVG 不受影响）")

    # ---- 发布前预检（文件层）
    try:
        from PIL import Image
        rgba = tmp / "alpha.png"
        Image.new("RGBA", (20, 20), (255, 255, 255, 255)).save(rgba, dpi=(600, 600))
        hits = pf.check_png(rgba, 300)
        record("预检：带透明通道的 PNG 应报错",
               any(h.code == "PFT-EXP-002" for h in hits), "；".join(map(str, hits)))
        nodpi = tmp / "nodpi.png"
        Image.new("RGB", (20, 20), (255, 255, 255)).save(nodpi)
        hits = pf.check_png(nodpi, 300)
        record("预检：没写 DPI 的 PNG 报「未检查」而不是通过",
               any(h.code == "PFT-EXP-001" and h.severity == "info" for h in hits),
               "；".join(map(str, hits)))
    except ImportError:
        record("预检：PNG 用例", True, "没有 Pillow，跳过")

    bad_uni = tmp / "uni.svg"
    bad_uni.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="120pt" height="40pt">'
                       '<text x="5" y="20" font-size="9">\ufffd \u63a8</text></svg>',
                       encoding="utf-8")
    hits = pf.check_svg(bad_uni, 300)
    record("预检：替换字符应报错",
           any(h.code == "PFT-EXP-004" for h in hits), "；".join(map(str, hits)))

    # 空白 PDF：字体、DPI 全都合格，页面上却什么都没有。
    # 多面板图导出时最容易出这个，而它一路通过所有文件层检查。
    try:
        import pymupdf as _pym
        blank = tmp / "blank.pdf"
        _d = _pym.open()
        _d.new_page(width=462, height=275)
        _d.save(str(blank))
        _d.close()
        bhits = pf.check_pdf(blank)
        record("预检：空白 PDF 应报错",
               any(h.code == "PFT-EXP-007" for h in bhits), "；".join(map(str, bhits))[:60])
    except ImportError:
        record("预检：空白 PDF 用例", True, "没有 PyMuPDF，跳过")

    # ---- 规则目录完整性：代码里引用的编号必须在目录里，且每条都写了局限
    cat = pf.load_rules()
    ids = {r["id"] for r in cat["rules"]}
    code_txt = "\n".join(p.read_text(encoding="utf-8") for p in HERE.glob("*.py"))
    used = set(re.findall(r"PFT-[A-Z]+-\d{3}", code_txt))
    record("规则目录：代码引用的编号都已登记", used <= ids,
           "未登记：" + ",".join(sorted(used - ids)))
    need = ("id", "scope", "title", "check", "source", "automated", "manual", "limit")
    missing = [r.get("id", "?") for r in cat["rules"] if not all(k in r for k in need)]
    record("规则目录：每条都有出处与局限（limit）", not missing, ",".join(missing))

    # ---- QC 报告：固定声明不构成合规认证
    rep = tmp / "qc.json"
    write_report(rep, [Finding("PFT-QC-001", "error", "测试用")], targets=["x"])
    d = json.loads(rep.read_text(encoding="utf-8"))
    record("QC 报告：声明 publication_ready_claim = false",
           d.get("publication_ready_claim") is False and d["summary"]["error"] == 1,
           json.dumps(d.get("summary"), ensure_ascii=False))

    # ---- 版式 profile
    bad_prof = []
    for pid in profiles.list_profiles():
        raw = json.loads((profiles.PROFILE_DIR / f"{pid}.json").read_text(encoding="utf-8"))
        bad_prof += [f"{pid}: {e}" for e in profiles.validate(raw)]
    record("profile：内置文件字段合法（含出处与状态）", not bad_prof, "；".join(bad_prof))

    import contextlib
    import io
    ok_all, failed = [], []
    with contextlib.redirect_stderr(io.StringIO()):     # template 状态的提醒不用刷屏
        for pid in profiles.list_profiles():
            try:
                make_table.build_docx(dict(TABLE_APA, profile=pid),
                                      tmp / f"profile_{pid}.docx")
                ok_all.append(pid)
            except SystemExit as e:
                failed.append(f"{pid}: {str(e).splitlines()[0]}")
    record("profile：每个内置 profile 都能生成", not failed, "；".join(failed))

    ok, msg = expect_exit(lambda: make_table.build_docx(
        {"columns": [{"key": "a"}], "rows": [], "style": {"fontsize": 9}},
        tmp / "bad_style.docx"), "未知字段")
    record("负向：style 里拼错字段应报错", ok, msg)

    # 同一份 docx，按不同 profile 判，结论应当不同——这是 profile 存在的理由
    zh_spec = {"number": "9", "title": "t", "profile": "zh-journal",
               "columns": [{"key": "v", "label": "变量"}, {"key": "r", "label": "r"}],
               "rows": [{"v": "信任", "r": "0.53"}]}
    zh_path = tmp / "zh_profile.docx"
    with contextlib.redirect_stderr(io.StringIO()):
        make_table.build_docx(zh_spec, zh_path)

    def has_lz(ws):
        return any(getattr(x, "code", "") == "PFT-TBL-003" for x in ws)

    _, w_def = check_figure.check_docx(zh_path)
    _, w_apa = check_figure.check_docx(zh_path, profiles.load("apa7", quiet=True))
    record("profile：中文刊保留前导零，不报", not has_lz(w_def), "；".join(map(str, w_def)))
    record("profile：同一文件按 APA 判则报前导零", has_lz(w_apa), "；".join(map(str, w_apa)))

    # ---- 审已有表：故意做一张坏表，看能不能逐类报出来
    bad = tmp / "bad_table.docx"
    _make_bad_table(bad)
    prof_apa = profiles.load("apa7", quiet=True)
    hits = audit_table.audit(bad, prof_apa)
    codes = {f.code for f in hits}
    want = {"PFT-TBL-009", "PFT-TBL-001", "PFT-TBL-002",
            "PFT-TBL-010", "PFT-TBL-011", "PFT-TBL-004"}
    record("审核器：坏表能逐类报出（表题/边框/字号/符号/小数位）",
           want <= codes, "报了：" + ",".join(sorted(codes)))

    good = tmp / "good_table.docx"
    make_table.build_docx(RICH_TABLE, good)
    hits2 = audit_table.audit(good, prof_apa)
    errs2 = [f for f in hits2 if f.severity == "error"]
    record("审核器：自己生成的表无硬伤", not errs2, "；".join(map(str, errs2)))

    # ---- 表型模板：每个示例规格都要能生成且无硬伤
    spec_dir = HERE.parent / "assets" / "table_specs"
    bad_specs = []
    for p in sorted(spec_dir.glob("*.json")):
        spec = json.loads(p.read_text(encoding="utf-8"))
        try:
            make_table.build_docx(spec, tmp / f"tmpl_{p.stem}.docx")
            errs_t, _ = check_figure.check_docx(tmp / f"tmpl_{p.stem}.docx")
            if errs_t:
                bad_specs.append(f"{p.name}: {errs_t[0]}")
        except SystemExit as e:
            bad_specs.append(f"{p.name}: {str(e).splitlines()[0]}")
    record("表型模板：示例规格都能生成且无硬伤", not bad_specs, "；".join(bad_specs))

    mm = json.loads((spec_dir / "table5_moderated_mediation.example.json")
                    .read_text(encoding="utf-8"))
    mm_text = json.dumps(mm, ensure_ascii=False)
    record("表型模板：有调节的中介含 index 与三个条件取值点",
           all(k in mm_text for k in ("Index", "−1 SD", "mean", "+1 SD",
                                      "Conditional indirect")),
           "缺：" + ",".join(k for k in ("Index", "−1 SD", "mean", "+1 SD",
                                         "Conditional indirect") if k not in mm_text))

    # ---- 合并单元格的宽度：跨列标题必须真的跨过去
    mm_out = tmp / "moderated_mediation.docx"
    make_table.build_docx(mm, mm_out)
    errs_mm, _ = check_figure.check_docx(mm_out)
    record("表型模板：跨列标题的单元格宽度与列网格一致",
           not any(getattr(f, "code", "") == "PFT-TBL-012" for f in errs_mm),
           "；".join(map(str, errs_mm)))

    from docx import Document
    from docx.oxml.ns import qn
    d_bad = Document(str(mm_out))
    tc_last = d_bad.tables[0].rows[0]._tr.findall(qn("w:tc"))[-1]
    tc_last.tcPr.find(qn("w:tcW")).set(qn("w:w"), "3456")     # 单列宽度的两倍
    bad_width = tmp / "bad_width.docx"
    d_bad.save(str(bad_width))
    errs_bw, _ = check_figure.check_docx(bad_width)
    record("负向：合并单元格宽度不对应报错",
           any(getattr(f, "code", "") == "PFT-TBL-012" for f in errs_bw),
           "；".join(map(str, errs_bw)))

    # ---- 多个跨列标题：第二个不能把第一个的宽度改回单列
    # 曾经的 bug：设宽度的循环写在了 spanner 循环**里面**，第二遍 table.cell(0, j)
    # 会把 j 映射到已合并的单元格上，于是第一个跨列标题又变回单列宽。
    two_sp = {
        "number": "8", "profile": "zh-journal", "title": "两个跨列标题",
        "spanners": [{"label": "信度与收敛效度", "span": ["al", "cr", "av"]},
                     {"label": "相关系数", "span": ["c1", "c2"]}],
        "columns": [{"key": "v", "label": "变量"},
                    {"key": "m", "label": "均值"},
                    {"key": "sd", "label": "标准差"},
                    {"key": "al", "label": "α"},
                    {"key": "cr", "label": "CR"},
                    {"key": "av", "label": "AVE"},
                    {"key": "c1", "label": "1"},
                    {"key": "c2", "label": "2"}],
        "rows": [{"v": "1. 品牌信任", "m": "4.21", "sd": "1.08", "al": "0.89",
                  "cr": "0.91", "av": "0.72", "c1": "0.85", "c2": ""}],
    }
    two_out = tmp / "two_spanners.docx"
    make_table.build_docx(two_sp, two_out)
    errs_two, _ = check_figure.check_docx(two_out)
    record("两个跨列标题：两块的宽度都正确",
           not any(getattr(f, "code", "") == "PFT-TBL-012" for f in errs_two),
           "；".join(map(str, errs_two)))
    record("中文按字断行：7 个汉字的最小宽度是单字宽",
           make_table.word_min_width("信度与收敛效度", 9.0) < 12.0,
           f"{make_table.word_min_width('信度与收敛效度', 9.0):.1f} pt")


    k_spec = {
        "number": "9", "title": "t", "profile": "apa7",
        "columns": [{"key": "v", "label": "Model"},
                    {"key": "aic", "label": "AIC", "decimals": 2,
                     "leading_zero": True, "thousands": True}],
        "rows": [{"v": "M1", "aic": "10109.33"}, {"v": "M2", "aic": "1234.5"}],
    }
    k_out = tmp / "thousands.docx"
    make_table.build_docx(k_spec, k_out)
    from docx import Document as _Doc
    kd = _Doc(str(k_out))
    vals = [kd.tables[0].cell(i, 1).text for i in (1, 2)]
    record("表格：千分位格式（10,109.33 / 1,234.50）",
           vals == ["10,109.33", "1,234.50"], str(vals))

    # ---- 多面板表：一套网格、每个面板自带表头
    PANEL_SPEC = {
        "number": "9", "profile": "apa7", "title": "Panel demo",
        "panels": [
            {"label": "A: First", "stub": "",
             "columns": [{"label": "X (1)", "span": 2}, {"label": "Y (2)", "span": 2}],
             "rows": [{"label": "row", "values": [".11", ".22"]}]},
            {"label": "B: Second", "stub": "",
             "columns": [{"label": "Z (3)", "span": 4}],
             "rows": [{"label": "r1", "values": [".33"]},
                      {"label": "r2", "values": [".44"]}]},
        ],
        "notes": {"general": "demo."},
    }
    panel_out = tmp / "panel.docx"
    make_table.build_docx(PANEL_SPEC, panel_out)
    from docx import Document as _Doc2
    pd = _Doc2(str(panel_out))
    pt = pd.tables[0]
    record("多面板表：两块面板共用一套网格（7 行 × 5 列）",
           len(pt.rows) == 7 and len(pt.columns) == 5,
           f"{len(pt.rows)} 行 × {len(pt.columns)} 列")
    errs_p, _ = check_figure.check_docx(panel_out)
    record("多面板表：自检通过", not errs_p, "；".join(map(str, errs_p)))

    # ---- 列宽必须放得下最长的词（否则 Word 从词中间断开）
    from docx.oxml.ns import qn as _qn
    from textmetrics import text_width as tw, word_min_width as wmw
    # 实测值：9 pt Times 下 "Relational" 38.0 pt、"Personalization" 54.2 pt
    record("字宽估算：Relational 误差 < 8%",
           abs(tw("Relational", 9.0) - 38.0) / 38.0 < 0.08,
           f"{tw('Relational', 9.0):.1f} pt vs 实测 38.0")
    record("字宽估算：Personalization 误差 < 10%",
           abs(tw("Personalization", 9.0) - 54.2) / 54.2 < 0.10,
           f"{tw('Personalization', 9.0):.1f} pt vs 实测 54.2")
    record("字宽估算：旧实现（0.55 em）会高估三成",
           tw("Relational", 9.0) < 0.55 * 9.0 * 10,
           f"{tw('Relational', 9.0):.1f} < {0.55 * 9 * 10:.1f}")
    record("最短词：按词而不是整行取",
           wmw("Customer Satisfaction (8)", 9.0) < tw("Customer Satisfaction (8)", 9.0))

    # 真实用例：6 列的面板每列跨 2 格、4 列的面板每列跨 3 格。
    # 均分会让前者的格子只有后者的 2/3 宽，"Relational" 就被断成 "Relationa / l"。
    WIDE_SPEC = {
        "number": "10", "profile": "apa7", "title": "Wide panel",
        "style": {"size": 9.0}, "table_width": 468,
        "panels": [
            {"label": "A: six", "stub": "",
             "columns": [{"label": f"{w} ({i})", "span": 2} for i, w in enumerate(
                 ["Affective", "Cognitive", "Physical", "Relational", "Sensorial", "Symbolic"], 1)],
             "rows": [{"label": "Personalization \u2192 CX Dimension",
                       "values": [".42***"] * 6}]},
            {"label": "B: four", "stub": "",
             "columns": [{"label": f"{w} ({i})", "span": 3} for i, w in enumerate(
                 ["Customer Attitude", "Customer Satisfaction", "Word of Mouth", "Customer Loyalty"], 7)],
             "rows": [{"label": "Affective \u2192 Consequence",
                       "values": [".34***", ".08a", ".26***", ".65***"]}]},
        ],
    }
    wide_out = tmp / "wide_panel.docx"
    make_table.build_docx(WIDE_SPEC, wide_out)
    werrs, _ = check_figure.check_docx(wide_out)
    record("多面板表：窄面板的列也放得下最长的词（不断词）", not werrs,
           "；".join(map(str, werrs)))
    wd = _Doc2(str(wide_out))
    grid_w = [int(g.get(_qn("w:w"))) / 20
              for g in wd.tables[0]._tbl.tblGrid.findall(_qn("w:gridCol"))]
    cell_a = sum(grid_w[1:3])
    record("多面板表：6 列面板的单元格宽度 ≥ Relational 的宽度",
           cell_a >= tw("Relational", 9.0) + 6, f"{cell_a:.1f} pt")

    # 负向：把跨两格的那个单元格压到 12 pt，必须报断词风险
    squeezed = tmp / "squeezed.docx"
    sd = _Doc2(str(wide_out))
    for tc in sd.tables[0]._tbl.iter(_qn("w:tc")):
        txt = "".join(n.text or "" for n in tc.iter(_qn("w:t")))
        if txt.strip().startswith("Relational"):
            tcW = tc.tcPr.find(_qn("w:tcW"))
            tcW.set(_qn("w:w"), "240")          # 12 pt
    sd.save(str(squeezed))
    serrs, _ = check_figure.check_docx(squeezed)
    record("负向：列被压窄 应报断词风险",
           any(getattr(e, "code", "") == "PFT-TBL-013" for e in serrs),
           "；".join(getattr(e, "code", "?") for e in serrs))

    # ---- 横向页面：过宽的表放在横向页上，版心 9 in
    LAND_SPEC = {
        "number": "11", "profile": "apa7", "title": "Wide table",
        "landscape": True,
        "columns": [{"key": f"c{i}", "label": f"Col {i}", "decimals": 2}
                    for i in range(1, 13)],
        "rows": [{"c1": f"Item {r}"} | {f"c{i}": f"0.{r}{i}" for i in range(2, 13)}
                 for r in range(1, 4)],
    }
    land_out = tmp / "landscape.docx"
    make_table.build_docx(LAND_SPEC, land_out)
    from docx.enum.section import WD_ORIENT
    ld = _Doc2(str(land_out))
    lsec = ld.sections[0]
    record("横向页：orientation 与页宽页高都对调了",
           lsec.orientation == WD_ORIENT.LANDSCAPE
           and abs(lsec.page_width.pt - 792) < 1 and abs(lsec.page_height.pt - 612) < 1,
           f"{lsec.orientation} {lsec.page_width.pt:.0f}×{lsec.page_height.pt:.0f}")
    land_grid = [int(g.get(_qn("w:w"))) / 20
                 for g in ld.tables[0]._tbl.tblGrid.findall(_qn("w:gridCol"))]
    record("横向页：不指定 table_width 时铺满版心（648 pt）",
           abs(sum(land_grid) - 648) < 1.5, f"{sum(land_grid):.1f} pt")
    lerrs, _ = check_figure.check_docx(land_out)
    record("横向页：自检通过", not lerrs, "；".join(map(str, lerrs)))

    # 负向：同样的表放在纵向页上，必须报「超出页宽」
    narrow = json.loads(json.dumps(LAND_SPEC))
    narrow["landscape"] = False
    narrow["table_width"] = 640
    bad_land = tmp / "landscape_on_portrait.docx"
    make_table.build_docx(narrow, bad_land)
    blerrs, _ = check_figure.check_docx(bad_land)
    record("负向：纵向页放过宽的表 应报 PFT-TBL-014",
           any(getattr(e, "code", "") == "PFT-TBL-014" for e in blerrs),
           "；".join(getattr(e, "code", "?") for e in blerrs))

    # A4 的版心要跟着 page.size 变，不能写死 6.5 in
    a4 = tmp / "a4.json"
    a4.write_text(json.dumps({
        "profile_id": "a4test", "label": "A4 测试", "status": "simulated",
        "page": {"size": "a4", "margins_in": [1.0, 1.0, 1.0, 1.0]},
    }, ensure_ascii=False), encoding="utf-8", newline="\n")
    a4p = profiles.load(str(a4), quiet=True)
    record("A4：纵向版心 451.5 pt、横向 697.5 pt",
           abs(profiles.text_width_pt(a4p) - 451.44) < 1
           and abs(profiles.text_width_pt(a4p, True) - 697.68) < 1,
           f"{profiles.text_width_pt(a4p):.1f} / {profiles.text_width_pt(a4p, True):.1f}")

    # ---- 一格多行：一个格子里按研究竖排四个值（宽表的常见写法）
    STACK_SPEC = {
        "number": "12", "profile": "apa7", "title": "Stacked", "landscape": True,
        "columns": [{"key": "v", "label": "Dimension"},
                    {"key": "s4", "label": "Study 4", "decimals": 2},
                    {"key": "a", "label": "A", "decimals": 2},
                    {"key": "c", "label": "C", "decimals": 2}],
        "rows": [
            {"v": "Affective", "s4": ".75", "a": ["0.75", ".9", ".88", ".86"], "c": ""},
            {"v": "Cognitive", "s4": ".55", "a": [".17", ".59", ".63", ".80"],
             "c": [".55", ".62", ".41", ".73"]},
        ],
        "notes": {"general": "demo."},
    }
    stack_out = tmp / "stacked.docx"
    make_table.build_docx(STACK_SPEC, stack_out)
    sd2 = _Doc2(str(stack_out))
    st2 = sd2.tables[0]
    cell_a = st2.rows[1].cells[2]
    record("一格多行：四个值写成四段（不是四个 run）",
           len(cell_a.paragraphs) == 4 and all(len(p.runs) == 1 for p in cell_a.paragraphs),
           f"{len(cell_a.paragraphs)} 段")
    record("一格多行：decimals 逐行生效（0.75→.75、.9→.90）",
           [p.text for p in cell_a.paragraphs] == [".75", ".90", ".88", ".86"],
           str([p.text for p in cell_a.paragraphs]))
    record("一格多行：全空的格子只写一段（三角矩阵的下半部）",
           len(st2.rows[1].cells[3].paragraphs) == 1,
           f"{len(st2.rows[1].cells[3].paragraphs)} 段")
    serrs2, _ = check_figure.check_docx(stack_out)
    record("一格多行：自检通过", not serrs2, "；".join(map(str, serrs2)))

    # 负向：把堆叠列压到 12 pt，必须报断词风险
    sq = _Doc2(str(stack_out))
    for tc in sq.tables[0]._tbl.iter(_qn("w:tc")):
        txt = "".join(n.text or "" for n in tc.iter(_qn("w:t")))
        if txt.startswith(".17"):
            tc.tcPr.find(_qn("w:tcW")).set(_qn("w:w"), "240")
    sq_out = tmp / "stacked_squeezed.docx"
    sq.save(str(sq_out))
    sqerrs, _ = check_figure.check_docx(sq_out)
    record("负向：堆叠格被压窄 应报断词风险",
           any(getattr(e, "code", "") == "PFT-TBL-013" for e in sqerrs),
           "；".join(getattr(e, "code", "?") for e in sqerrs))

    # check_numbers 必须按段落分隔取数字，否则一格四行会粘成 "0.750.90"
    import check_numbers as cn
    toks, _neg = cn._tokens("0.75 0.90 0.88 0.86")
    record("数字核对：段间补空格后不会粘出假数字",
           toks["0.75"] == 1 and "0.750" not in toks,
           str(sorted(toks)))

    # ---- 多系列折线：三条线 + 置信带（Figure 2 那种）
    LINE_SPEC = {
        "xlabel": "Stage", "ylabel": "CX (1-7)", "y_range": [1, 7],
        "x_ticks": [{"v": 1, "label": "Pre"}, {"v": 2, "label": "During"},
                    {"v": 3, "label": "Post"}],
        "markers": True, "note": "Note. Bands are 95% CI.",
        "series": [
            {"label": "Cognitive", "x": [1, 2, 3], "y": [4.8, 4.5, 4.9],
             "lo": [4.7, 4.4, 4.8], "hi": [4.9, 4.6, 5.0]},
            {"label": "Affective", "x": [1, 2, 3], "y": [4.3, 4.8, 4.6],
             "lo": [4.2, 4.7, 4.5], "hi": [4.4, 4.9, 4.7]},
            {"label": "Behavioral", "x": [1, 2, 3], "y": [4.1, 4.2, 4.0],
             "lo": [4.0, 4.1, 3.9], "hi": [4.2, 4.3, 4.1]},
        ],
    }
    line_svg, line_warns = line_plot.build(LINE_SPEC)
    line_out = line_svg.save(tmp / "line.svg")
    record("折线图：三条系列 + 置信带生成成功", line_out.exists())
    record("折线图：规格无警告", not line_warns, "；".join(line_warns))
    lerrs, _ = check_figure.check_svg(line_out)
    record("折线图：图检查通过", not lerrs, "；".join(map(str, lerrs)))
    ltxt = line_out.read_text(encoding="utf-8")
    record("折线图：三条线用三种线型（黑白可辨）",
           ltxt.count('stroke-dasharray') >= 2, f"{ltxt.count('stroke-dasharray')} 处虚线")

    # 线端标签不该叠在一起
    ys = [float(m) for m in re.findall(r'<text x="[\d.]+" y="([\d.]+)"[^>]*>(?:Cognitive|Affective|Behavioral)</text>', ltxt)]
    record("折线图：线端标签按最小间距推开",
           len(ys) == 3 and all(abs(ys[i + 1] - ys[i]) >= 9.4 for i in range(len(ys) - 1)),
           str([round(y, 1) for y in ys]))

    bad_line = json.loads(json.dumps(LINE_SPEC))
    bad_line["series"][0].pop("hi")
    record("负向：置信带只给一半 应报错",
           any("lo 或 hi" in e for e in line_plot.validate(bad_line)),
           "；".join(line_plot.validate(bad_line))[:60])

    likert = json.loads(json.dumps(LINE_SPEC))
    likert.pop("y_range")
    record("负向：李克特数据没给 y_range 应提示",
           any("李克特" in e for e in line_plot.validate(likert)),
           "；".join(line_plot.validate(likert))[:60])

    # ---- 顶部图例：图例里的样例线型必须与图上一致
    top_spec = json.loads(json.dumps(LINE_SPEC))
    top_spec["legend"] = "top"
    top_spec["note"] = ("Note. 这是一条故意写得很长的图注，用来确认它会被折行而不是"
                        "被画布右边缘裁掉——被裁掉的文字在屏幕上完全看不出来。")
    top_svg, _ = line_plot.build(top_spec)
    top_out = top_svg.save(tmp / "line_top.svg")
    ttxt = top_out.read_text(encoding="utf-8")
    import collections as _c
    dash_counts = _c.Counter(re.findall(r'stroke-dasharray="([^"]+)"', ttxt))
    record("折线图：图例的样例线型与图上一致（每种虚线各出现两次）",
           dash_counts and all(v == 2 for v in dash_counts.values()),
           str(dict(dash_counts)))
    note_texts = re.findall(r'<text[^>]*>([^<]*被画布[^<]*)</text>', ttxt)
    record("折线图：超长图注折行而不是被裁掉",
           len(re.findall(r'<text[^>]*font-size="8.2"[^>]*>', ttxt)) >= 2,
           f"{len(re.findall(chr(39) + '<text[^>]*font-size=.8.2.[^>]*>' + chr(39), ttxt))} 段")
    tall, short = top_svg.h, line_svg.h
    record("折线图：顶部图例占用上方留白（画布变高）", tall > short,
           f"{short:.0f} → {tall:.0f} pt")

    # 一行放不下时图例折成多行
    wide_labels = json.loads(json.dumps(top_spec))
    for i, s in enumerate(wide_labels["series"]):
        s["label"] = f"A very long series label number {i + 1}"
    wide_svg, _ = line_plot.build(wide_labels)
    wide_out = wide_svg.save(tmp / "line_top2.svg")
    wtxt = wide_out.read_text(encoding="utf-8")
    rows_y = set(re.findall(r'<text x="[\d.]+" y="([\d.]+)"[^>]*>A very long', wtxt))
    record("折线图：图例一行放不下时折成多行", len(rows_y) >= 2, f"{len(rows_y)} 行")

    # ---- 审别人的图：一个「学生常见写法」的 Excel 图（饼图 + 网格线 + 图例在右 + 无轴标题）
    import zipfile
    bad_xlsx = tmp / "bad_chart.xlsx"
    with zipfile.ZipFile(bad_xlsx, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/worksheets/sheet1.xml",
                   '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                   '<sheetData>'
                   '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c></row>'
                   '<row r="2"><c r="A2" t="s"><v>2</v></c><c r="B2"><v>4.2</v></c></row>'
                   '<row r="3"><c r="A3" t="s"><v>3</v></c><c r="B3"><v>4.8</v></c></row>'
                   "</sheetData></worksheet>")
        z.writestr("xl/sharedStrings.xml",
                   '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                   "<si><t>Stage</t></si><si><t>Cognitive CX</t></si>"
                   "<si><t>Pre</t></si><si><t>Post</t></si></sst>")
        z.writestr("xl/charts/chart1.xml",
                   '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
                   'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:chart>'
                   "<c:plotArea><c:pieChart><c:ser>"
                   "<c:tx><c:strRef><c:f>Sheet1!$B$1</c:f></c:strRef></c:tx>"
                   "<c:cat><c:strRef><c:f>Sheet1!$A$2:$A$3</c:f></c:strRef></c:cat>"
                   "<c:val><c:numRef><c:f>Sheet1!$B$2:$B$3</c:f></c:numRef></c:val>"
                   "</c:ser></c:pieChart>"
                   "<c:catAx><c:majorGridlines/></c:catAx><c:valAx/></c:plotArea>"
                   '<c:legend><c:legendPos val="r"/></c:legend></c:chart></c:chartSpace>')
    aprof = profiles.load("apa7", quiet=True)
    afind, aspec = audit_figure.audit_xlsx(bad_xlsx, aprof)
    acodes = {f.code for f in afind}
    record("审图：饼图被判为必须改",
           any(f.code == "PFT-AUD-007" and f.severity == "error" for f in afind), str(sorted(acodes)))
    record("审图：网格线、图例在右、缺轴标题都报出",
           {"PFT-AUD-009", "PFT-AUD-010", "PFT-AUD-011"} <= acodes, str(sorted(acodes)))
    record("审图：能从图表引用反查出数据（可重画）",
           bool(aspec) and aspec["series"][0]["label"] == "Cognitive CX"
           and aspec["series"][0]["y"] == [4.2, 4.8],
           json.dumps(aspec["series"][0] if aspec else {}, ensure_ascii=False)[:70])
    record("审图：分类轴转成 x_ticks",
           bool(aspec) and [t["label"] for t in aspec.get("x_ticks", [])] == ["Pre", "Post"],
           str(aspec.get("x_ticks") if aspec else None))

    # 负向：不是 xlsx 的文件要说清楚
    notzip = tmp / "fake.xlsx"
    notzip.write_bytes(b"not a zip")
    nfind, nspec = audit_figure.audit_xlsx(notzip, aprof)
    record("负向：坏 xlsx 应报错而不是崩",
           nspec is None and any(f.severity == "error" for f in nfind),
           "；".join(map(str, nfind))[:50])

    # ---- 负零与连字符占位（这两条是从一篇真实论文的表里发现的）
    nz_spec = {
        "number": "9", "profile": "zh-journal", "title": "负零",
        "columns": [{"key": "h", "label": "假设"},
                    {"key": "b", "label": "路径值（t 值）"}],
        "rows": [{"h": "H14", "b": "-0.00（-0.08）"},
                 {"h": "H15", "b": "----"},
                 {"h": "H16", "b": "-0.004（-0.31）"}],
    }
    nz_out = tmp / "negzero.docx"
    make_table.build_docx(nz_spec, nz_out)
    nz_errs, nz_warns = check_figure.check_docx(nz_out)
    nz_all = list(nz_errs) + list(nz_warns)     # ---- 是提示级，不能只看 errs
    nz_codes = {getattr(f, "code", "") for f in nz_all}
    record("负零：-0.00 即使嵌在括号前也要报",
           "PFT-TBL-015" in nz_codes, str(sorted(nz_codes)))
    record("负零：-0.004 不算负零",
           sum(1 for f in nz_errs if getattr(f, "code", "") == "PFT-TBL-015") == 1,
           f"{sum(1 for f in nz_errs if getattr(f, 'code', '') == 'PFT-TBL-015')} 条")
    record("占位符：---- 应提示改用破折号",
           "PFT-TBL-016" in nz_codes, str(sorted(nz_codes)))

    # ---- 审表不吃非 docx：要说清楚，不能崩在 python-docx 里
    notdocx = tmp / "notword.pdf"
    notdocx.write_bytes(b"%PDF-1.4\n%%EOF\n")
    tf = audit_table.audit(notdocx, profiles.load("apa7", quiet=True))
    record("审表：给 PDF 时应报「只能读 docx」而不是崩",
           any(getattr(f, "code", "") == "PFT-TBL-000" for f in tf),
           "；".join(map(str, tf))[:50])

    # ---- 审 PDF 里的图：逐页 + 有效分辨率
    try:
        import pymupdf as _pym
        doc2 = _pym.open()
        doc2.new_page(width=400, height=300)          # 第 1 页空白
        p2 = doc2.new_page(width=400, height=300)
        tiny = tmp / "tiny.png"
        from PIL import Image as _Img
        _Img.new("RGB", (40, 30), (200, 60, 60)).save(tiny)
        p2.insert_image(_pym.Rect(50, 50, 350, 275), filename=str(tiny))
        pdf2 = tmp / "paper.pdf"
        doc2.save(str(pdf2))
        doc2.close()
        f2 = audit_figure.audit_pdf(pdf2, profiles.load("apa7", quiet=True))
        codes2 = {getattr(f, "code", "") for f in f2}
        record("审图 PDF：空白页要报出来（不能只看第 1 页）",
               any(getattr(f, "code", "") == "PFT-AUD-015" and "1" in str(f) for f in f2),
               "；".join(map(str, f2))[:60])
        record("审图 PDF：低分辨率嵌入图要报出来",
               "PFT-AUD-016" in codes2, str(sorted(codes2)))
    except ImportError:
        record("审图 PDF 用例", True, "没有 PyMuPDF 或 Pillow，跳过")

    bad_panel = json.loads(json.dumps(PANEL_SPEC))
    bad_panel["panels"][1]["columns"][0]["span"] = 3     # 跨度合计 3 ≠ 4
    ok, msg = expect_exit(lambda: make_table.build_docx(bad_panel, tmp / "x.docx"), "跨度")
    record("负向：面板跨度不相等应报错", ok, msg)

    # ---- 数字与来源核对
    import check_numbers as cn
    num_spec = {"number": "9", "title": "t", "profile": "apa7",
                "columns": [{"key": "a", "label": "A"},
                            {"key": "b", "label": "B", "decimals": 2, "leading_zero": True}],
                "rows": [{"a": "x", "b": "331.01"}, {"a": "y", "b": "298.86"}]}
    num_out = tmp / "num.docx"
    make_table.build_docx(num_spec, num_out)
    num_src = tmp / "num_src.txt"
    num_src.write_text("Brand 195 331.01 120 .94\nEmployee 203 298.86 120 .94\n",
                       encoding="utf-8")
    mine_n, _ = cn.docx_numbers(num_out)
    src_n, _ = cn.source_numbers(num_src)
    record("数字核对：表内数字都能在来源中找到",
           not cn.compare(mine_n, src_n), "；".join(map(str, cn.compare(mine_n, src_n))))

    d_num = _Doc2(str(num_out))
    d_num.tables[0].cell(1, 1).paragraphs[0].runs[0].text = "331.10"   # 改错一位
    d_num.save(str(num_out))
    mine_bad, _ = cn.docx_numbers(num_out)
    record("负向：改错一位数应被数字核对抓到",
           bool(cn.compare(mine_bad, src_n)), "没有报出来")

    # ---- 负向用例：必须报错，且说清错在哪
    bad_specs = {
        "id 打错": ({**DIAG_CONCEPT, "edges": [{"from": "a", "to": "zzz"}]}, "不在 nodes 里"),
        "缺 from": ({**DIAG_CONCEPT, "edges": [{"to": "t"}]}, "缺 `from`"),
        "重复 id": ({"nodes": [{"id": "a"}, {"id": "a"}], "edges": []}, "重复"),
        "未知键": ({"nodes": [{"id": "a", "lable": "A"}], "edges": []}, "未知键"),
        "调节变量只给一个端点": ({**DIAG_CONCEPT, "moderators": [{"id": "m", "on": ["t"]}]},
                                 "两个节点 id"),
    }
    for label, (spec, needle) in bad_specs.items():
        ok, msg = expect_exit(lambda s=spec: make_diagram.build(s), needle)
        record(f"负向：{label} 应报错", ok, msg)

    # ---- 负向用例：测量模型漏写 bidirectional 应警告
    spec = json.loads(json.dumps(DIAG_MEAS))
    for e in spec["edges"]:
        e.pop("bidirectional", None)
    _, w = make_diagram.build(spec)
    hit = any("双向箭头" in x for x in w)
    record("负向：测量模型漏写 bidirectional 应警告", hit, "；".join(w) or "没有警告")

    # ---- 确定性
    a, b2 = tmp / "d1.svg", tmp / "d2.svg"
    make_diagram.build(DIAG_CONCEPT)[0].save(str(a))
    make_diagram.build(DIAG_CONCEPT)[0].save(str(b2))
    record("确定性：同一输入两次生成的 SVG 一致", md5(a) == md5(b2))


def connection_problems(geo: dict) -> list[str]:
    """用 check_connections.py 的 CLI 复查几何快照，返回不合格项。

    走子进程而不是 import：这样顺带验证了脚本能被正常调用，
    也避免它读 sys.argv 时和自检脚本打架。
    """
    p = Path(tempfile.gettempdir()) / "_pft_selftest.geo.json"
    p.write_text(json.dumps(geo, ensure_ascii=False), encoding="utf-8")
    r = subprocess.run([sys.executable, str(HERE / "check_connections.py"), str(p)],
                       capture_output=True, text=True, encoding="utf-8")
    return [ln.strip() for ln in r.stdout.splitlines() if "✗" in ln]


def main() -> int:
    print(f"技能自检：{HERE}")
    with tempfile.TemporaryDirectory(prefix="pft-selftest-") as d:
        run(Path(d))
    bad = [r for r in RESULTS if not r[1]]
    print(f"\n{len(RESULTS) - len(bad)} / {len(RESULTS)} 通过")
    if bad:
        print("失败项：")
        for name, _, note in bad:
            print(f"  ✗ {name}  {note}")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
