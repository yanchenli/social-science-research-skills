#!/usr/bin/env python3
"""把论文变量提取结果（JSON）导出为 Excel 变量表（codebook）。

用法:
    python3 export_codebook.py <input.json> <output.xlsx>
    python3 export_codebook.py <input.json> --validate      # 只做预检，不生成文件
    python3 export_codebook.py <input.json> --mermaid       # 只打印 Mermaid 研究地图源码

Mermaid 研究地图:
    生成 xlsx 时会同步写出同名 .mmd 文件，并在 Studies 表追加 `mermaid_node` 列
    （每行一段节点定义，可直接复制拼接）。完整源码也会打印到终端，便于粘进报告。

JSON 结构:
{
  "paper": {"title": "", "authors": "", "year": "", "journal": "", "doi": ""},
  "studies":    [ {...} ],   # 字段见 references/excel-codebook.md
  "variables":  [ {...} ],
  "items":      [ {...} ],   # 量表逐题项（自陈测量）
  "indicators": [ {...} ],   # 二手数据「多指标合成变量」的构成指标（与题项语义不同）
  "hypotheses": [ {...} ],
  "stimuli":    [ {...} ],
  "effectsizes":[ {...} ]    # 元分析专属：一行 = 一个「对比基准 × 结果变量 × 模型」的效应量
}

行为:
  - 固定列顺序与列名，保证跨论文一致（元分析编码友好）
  - JSON 中缺失字段自动补 "NA"，未知字段追加到末尾
  - 表头加粗 + 底色 + 冻结首行 + 自动列宽
  - 解析失败时定位到具体行号/列号并给出常见原因提示，避免手工排查引号
"""

import json
import re
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
HEADER_FONT = Font(bold=True)
BODY_ALIGN = Alignment(vertical="top", wrap_text=True)

SHEETS = ["Studies", "Variables", "Items", "Indicators", "Hypotheses", "Stimuli",
          "EffectSizes"]

# 固定列顺序：与 references/excel-codebook.md 保持一致
SCHEMA = {
    "Studies": [
        "study_id", "study_label", "method_type", "n", "n_unit", "sample_source", "design",
        "data_structure", "time_span", "level_1_unit", "level_2_unit", "n_observations",
        "estimator",
        # —— 计算测量专用（其余方法填 NA）——
        "train_n", "val_n", "holdout_n", "model_spec", "model_metrics",
        # —— 分析模型专用（其余方法填 NA）——
        "model_setup", "key_assumptions", "solution_concept", "model_results",
        "main_purpose", "role_in_paper", "key_finding",
        "relation_to_other_studies", "sampling_strategy", "field_sites",
        "researcher_positionality", "data_volume", "data_collection_period",
        # —— 元分析专用（其余方法填 NA）——
        "search_databases", "search_terms", "additional_search", "inclusion_criteria",
        "n_primary_papers", "n_primary_studies", "k_effect_sizes",
        "effect_size_metric", "effect_size_conversion", "coding_reliability",
        "publication_bias_tests", "missing_data_handling", "software",
        "preregistration", "supplement_location", "notes", "mermaid_node",
    ],
    "Variables": [
        "study_id", "study_label", "method_type", "variable_name_en", "variable_name_cn",
        "role", "variable_type", "parent_theme", "theme_level",
        "conceptual_definition_en", "conceptual_definition_cn",
        "operationalization", "composite", "n_indicators", "composite_method",
        "measurement_model", "scale_type", "coding_direction",
        # items_given：论文实际给出的题项条数（常少于 n_items，如只给样例题项）
        "n_items", "items_given", "items_text", "response_scale", "scale_anchors",
        "reverse_coded",
        # 信度拆列：论文常同时报 CR / α / AVE / 平均载荷，一列自由文本装不下
        "reliability_type", "reliability_value",
        "cr", "alpha", "ave", "avg_loading",
        # 计算测量：模型在该构念上的表现，与外部基准的一致性
        "model_performance", "validation_agreement",
        "source_apa",
        "adaptation", "manipulation_conditions", "manipulation_materials", "procedure",
        "manipulation_check", "data_source", "coding_rule", "aggregation", "sample_n",
        # —— 元分析：调节变量的编码取值来源与水平 ——
        "coding_source", "coding_values", "k_studies",
        "notes", "missing_fields",
    ],
    "Items": [
        "study_id", "variable_name_en", "item_no", "item_en", "item_cn",
        "dimension", "reverse_coded", "response_scale",
        "factor_loading", "dimension_cr", "dimension_ave",
    ],
    # Indicators：二手数据中「多指标合成变量」的构成指标（与 Items 的题项语义不同，分开存放）
    "Indicators": [
        "study_id", "variable_name_en", "indicator_no", "indicator_name_en",
        "indicator_name_cn", "data_source_detail", "unit", "coding_direction",
        "standardized", "coverage", "factor_loading", "notes",
    ],
    "Hypotheses": [
        "study_id", "hypothesis_id", "hypothesis_en", "hypothesis_cn", "iv", "dv",
        "mediator", "moderator", "test_type", "test_method", "statistic", "p_value",
        "effect_size", "zero_order_r", "direction", "supported", "note",
    ],
    "Stimuli": [
        "study_id", "stimulus_set", "stimulus_type", "content_en", "content_cn",
        "quantity", "presentation_mode", "duration", "pretest", "source",
    ],
    # EffectSizes：元分析的核心产物（一行 = 一个「对比基准 × 结果变量 × 模型」的效应量）
    "EffectSizes": [
        "study_id", "analysis_id", "comparison", "outcome", "model",
        "k_studies", "k_effect_sizes", "total_n", "effect_size_metric",
        "effect_size_value", "se", "ci_low", "ci_high", "p_value",
        "heterogeneity_i2", "cles", "fail_safe_n", "note",
    ],
}

# 设计上就该留空的字段（留空有意义，不是漏填）
OPTIONAL_BLANK = {"missing_fields"}

# 受控词表（仅提示，不强制）
CONTROLLED_VOCAB = {
    "method_type": {"Experiment", "Survey", "Secondary Data", "Qualitative", "Mixed",
                    "Scale Development", "Expert Judgment", "Item Sorting",
                    "Literature Review", "Simulation", "Meta-analysis",
                    "Computational Measurement", "Machine Learning",
                    "Analytical Model"},
    "role": {"IV", "DV", "Mediator", "Moderator", "Control", "Manipulation Check",
             "Confound Check", "Alternative Explanation", "Theme", "Target Construct",
             "Context Factor", "Covariate", "Comparison Construct",
             # 元分析
             "Effect Size", "Conceptual Pillar"},
    "variable_type": {"Self-report", "Behavioral", "Objective", "Manipulated", "Coded",
                      "Expert-rated", "Binary", "Composite", "Theme",
                      # 元分析：由编码/外部数据/自动文本分析得到的研究层面特征
                      "Coded from literature", "Secondary index", "Automated text",
                      "Survey-rated", "Expert-rated average", "Derived effect size",
                      # 计算测量：自动图像识别 / 模型预测值
                      "Automated image", "Model-predicted"},
    "supported": {"Supported", "Partially supported", "Not supported", "Reversed",
                  "Observed", "Partially observed", "Not observed",
                  # 分析模型：命题由推导得到，不是被观察到
                  "Derived", "NA"},
    "reverse_coded": {"Yes", "No", "Partial", "NA"},
    "composite": {"Yes", "No", "NA"},
    "scale_type": {"Formative", "Reflective", "Single-item", "NA"},
    "theme_level": {"Core construct", "Marker", "Aggregate dimension",
                    "Second-order theme", "First-order concept", "Code",
                    "Conceptual pillar", "NA"},
    "data_structure": {"Cross-sectional", "Panel", "Time-series", "Nested",
                       "Cross-classified", "Repeated cross-section",
                       "Effect sizes nested in studies", "NA"},
    "test_type": {"Primary", "Replication", "Robustness", "Supplemental",
                  "Manipulation check", "Pilot", "Objective",
                  # 分析模型：由假设推导命题
                  "Derivation", "NA"},
    # 元分析中同一个效应量常在不同模型下各报一次，必须标明是哪一种
    "model": {"Unconditional", "Full meta-regression", "Subgroup analysis",
              "Bivariate moderator synthesis", "Sensitivity", "Publication bias", "NA"},
    "design": {"Between-subjects", "Within-subjects", "Mixed", "Factorial",
               "Double randomization", "Experimental-causal-chain",
               "Moderation-of-process", "Blocking design", "Field experiment",
               "Single-group", "Correlational", "Path analysis", "Multilevel",
               "Integrative review", "Systematic review", "Qualitative coding",
               "Theory building", "Expert rating", "Item sorting",
               "Cross-sectional",
               # 质性设计
               "Ethnography", "Netnography", "Participant observation",
               "Interview study", "Focus group", "Case study", "Grounded theory",
               "Phenomenology", "Content analysis", "Thematic analysis",
               "Discourse analysis", "Narrative analysis",
               # 元分析
               "Meta-analysis", "Meta-regression", "Subgroup analysis",
               "Bivariate moderator synthesis",
               # 仿真与计算模型（method_type = Simulation 时的对应设计）
               "Simulation", "Agent-based model",
               # 分析模型（method_type = Analytical Model 时的对应设计）
               "Analytical",
               # 二手数据的识别设计（此前缺失，必然误报）
               "Difference-in-differences", "Instrumental variable",
               "Regression discontinuity", "Event study", "Quasi-experiment",
               "NA"},
}

# ---------------------------------------------------------------- Mermaid 研究地图
# 方法 → (样式类名, 填充色, 边框色, 中文标签)
METHOD_STYLE = {
    "Experiment":         ("exp",  "#FCE4EC", "#C2185B", "实验"),
    "Survey":             ("surv", "#E3F2FD", "#1565C0", "问卷"),
    "Secondary Data":     ("arch", "#E8F5E9", "#2E7D32", "二手数据"),
    "Qualitative":        ("qual", "#FFF3E0", "#E65100", "质性"),
    "Literature Review":  ("rev",  "#F3E5F5", "#6A1B9A", "文献综述"),
    "Scale Development":  ("dev",  "#E0F7FA", "#00838F", "量表开发"),
    "Expert Judgment":    ("dev",  "#E0F7FA", "#00838F", "专家评定"),
    "Item Sorting":       ("dev",  "#E0F7FA", "#00838F", "条目分类"),
    "Mixed":              ("mix",  "#ECEFF1", "#455A64", "混合"),
    "Simulation":         ("mix",  "#ECEFF1", "#455A64", "仿真"),
    "Meta-analysis":      ("meta", "#FFF9C4", "#F9A825", "元分析"),
    "Computational Measurement": ("comp", "#EDE7F6", "#4527A0", "计算测量"),
    "Analytical Model":   ("ana",  "#FFF8E1", "#F57F17", "分析模型"),
}
DEFAULT_STYLE = ("mix", "#ECEFF1", "#455A64", "其他")

# 节点内各行的字数上限（太长会把图撑爆）
LABEL_MAX = 20
ROLE_MAX = 16
TIP_MAX = 140


def mm_safe(text, limit=None):
    """清洗文本，保证塞进 Mermaid 节点/ tooltip 不会让渲染崩。"""
    s = str(text or "").strip()
    s = s.replace('"', "'").replace("\n", " ").replace("\r", " ")
    s = re.sub(r"\s+", " ", s)
    if limit and len(s) > limit:
        s = s[: limit - 1].rstrip() + "…"
    return s


def mm_n(n):
    """样本量格式化：加千分位，非数字原样返回。"""
    s = mm_safe(n, 14)
    if s.isdigit():
        return f"{int(s):,}"
    return s


def mm_node_for(study):
    """为一个研究生成 Mermaid 节点定义（单行文本）。

    主行放「角色」（这个研究在论文里干什么），第二行 n + 方法，第三行原文标签。
    ——读图的人第一眼要的是推进逻辑，不是原文标题。
    """
    sid = mm_safe(study.get("study_id", ""))
    label = mm_safe(study.get("study_label", ""), LABEL_MAX)
    method = study.get("method_type", "") or ""
    method_cn = METHOD_STYLE.get(method, DEFAULT_STYLE)[3]
    n = mm_n(study.get("n", ""))
    role = mm_safe(study.get("role_in_paper", ""), ROLE_MAX)

    head = f"{sid} · {role}" if role and role != "NA" else sid
    second = f"n={n} · {method_cn}" if n and n != "NA" else method_cn
    lines = [head, second]
    if label and label != "NA" and label not in head:
        lines.append(label)
    return f'{sid}["' + "<br/>".join(lines) + '"]'


def mm_tooltip_for(study):
    """生成悬停 tooltip 文本：优先主要发现，其次主要目标。"""
    finding = mm_safe(study.get("key_finding", ""), TIP_MAX)
    purpose = mm_safe(study.get("main_purpose", ""), TIP_MAX)
    if finding and finding != "NA":
        return "主要发现：" + finding
    if purpose and purpose != "NA":
        return "目标：" + purpose
    return ""


def build_mermaid(studies, edges=None, groups=None, with_tooltip=True):
    """生成 Mermaid 流程图源码。

    edges  : 可选，边定义列表，如 ["S1A --> S1B", "S1A -.-> S1C"]。
             不提供时按 studies 顺序串联，此时务必人工核对推进逻辑。
    groups : 可选，{分组标题: [study_id, ...]}。不提供且节点 > 9 时按方法自动分组。
    """
    rows = [s for s in studies if isinstance(s, dict) and s.get("study_id")]
    if not rows:
        return ""
    if edges is not None and not isinstance(edges, list):
        edges = None
    if groups is not None and not isinstance(groups, dict):
        groups = None

    out = ["```mermaid", "flowchart LR"]
    by_id = {mm_safe(s.get("study_id", "")): s for s in rows}

    if groups:
        covered = set()
        for gi, (title, members) in enumerate(groups.items(), start=1):
            out.append(f'    subgraph G{gi}["{mm_safe(title, 20)}"]')
            for sid in members:
                s = by_id.get(mm_safe(sid))
                if s:
                    covered.add(mm_safe(sid))
                    out.append("        " + mm_node_for(s))
            out.append("    end")
        # 自定义分组容易漏掉 pretest 之类的辅助节点，补出来，避免边指向空节点
        orphans = [s for s in rows if mm_safe(s.get("study_id", "")) not in covered]
        if orphans:
            out.append("    %% 以下节点未列入 mermaid_groups，已自动补出")
            for s in orphans:
                out.append("    " + mm_node_for(s))
    else:
        auto = {}
        for s in rows:
            auto.setdefault(s.get("method_type", "") or "", []).append(s)
        use_group = len(rows) > 9 and 2 <= len(auto) <= 6
        if use_group:
            for gi, (method, members) in enumerate(auto.items(), start=1):
                cn = METHOD_STYLE.get(method, DEFAULT_STYLE)[3]
                out.append(f'    subgraph G{gi}["{gi}. {cn}"]')
                for s in members:
                    out.append("        " + mm_node_for(s))
                out.append("    end")
        else:
            for s in rows:
                out.append("    " + mm_node_for(s))

    ids = [mm_safe(s.get("study_id", "")) for s in rows]
    out.append("")
    if edges:
        for e in edges:
            out.append("    " + mm_safe(e, 200))
    else:
        out.append("    " + " --> ".join(ids))

    used = {}
    for s in rows:
        method = s.get("method_type", "") or ""
        used.setdefault(method, []).append(mm_safe(s.get("study_id", "")))
    out.append("")
    styles = {}
    for method in used:
        cls, fill, stroke, _ = METHOD_STYLE.get(method, DEFAULT_STYLE)
        styles.setdefault(cls, (fill, stroke))
    for cls, (fill, stroke) in styles.items():
        out.append(f"    classDef {cls} fill:{fill},stroke:{stroke},color:#000")
    for method, members in used.items():
        cls = METHOD_STYLE.get(method, DEFAULT_STYLE)[0]
        out.append(f"    class {','.join(members)} {cls}")

    if with_tooltip:
        tips = [(sid, mm_tooltip_for(s)) for sid, s in zip(ids, rows)]
        tips = [(sid, t) for sid, t in tips if t]
        if tips:
            out.append("")
            for sid, tip in tips:
                out.append(f'    click {sid} href "#" "{tip}"')

    out.append("```")
    return "\n".join(out)


# sheet 小写键 → 正式表名（提示信息里用，避免 capitalize 把 EffectSizes 打成 Effectsizes）
SHEET_LABEL = {s.lower(): s for s in SHEETS}

# 哪个 sheet 的哪些字段需要做受控词表检查
SHEET_VOCAB = {
    "variables": ["variable_type", "role", "reverse_coded", "composite",
                  "theme_level", "scale_type"],
    "studies": ["method_type", "data_structure", "design"],
    "hypotheses": ["supported", "test_type"],
    "effectsizes": ["model"],
}


# ---------------------------------------------------------------- JSON 预检
def load_json_with_diagnostics(path):
    """解析 JSON；失败时给出行号、列号、上下文与常见原因。"""
    try:
        with open(path, "r", encoding="utf-8") as f:
            text = f.read()
    except FileNotFoundError:
        die(f"找不到文件: {path}")
    except UnicodeDecodeError as e:
        die(f"文件不是 UTF-8 编码: {e}\n  解决: 用 UTF-8 重新保存（不要带 BOM）。")

    try:
        return json.loads(text), text
    except json.JSONDecodeError as e:
        lines = text.split("\n")
        ln, col = e.lineno, e.colno
        print("=" * 72)
        print(f"JSON 解析失败 —— 第 {ln} 行，第 {col} 列")
        print(f"错误信息: {e.msg}")
        print("-" * 72)
        start = max(1, ln - 2)
        end = min(len(lines), ln + 2)
        for i in range(start, end + 1):
            mark = ">>" if i == ln else "  "
            content = lines[i - 1]
            print(f"{mark} {i:>5} | {content}")
            if i == ln:
                print(f"       {' ' * (col - 1)}^ 出错位置")
        print("-" * 72)
        for hint in diagnose(e, lines[ln - 1] if ln <= len(lines) else ""):
            print(f"  · {hint}")
        print("=" * 72)
        sys.exit(1)


def diagnose(err, bad_line):
    """根据错误类型给出可操作的修复建议。"""
    hints = []
    msg = err.msg

    if "Expecting ',' delimiter" in msg or "Expecting property name" in msg:
        hints.append("最可能的原因：字符串内部出现了未转义的双引号 \" "
                     "（例如中文用了直角引号之外的 \"...\"，或复制英文原文时带引号）。")
        hints.append("修复：把字符串内部的 \" 换成中文引号「」，或写成 \\\" 。")
    if "Unterminated string" in msg:
        hints.append("字符串没有闭合：某一行的双引号数量是奇数。")
    if "Expecting value" in msg:
        hints.append("某处缺少值、多了尾随逗号，或用了单引号 ' 包裹字符串（JSON 只认双引号）。")

    n = bad_line.count('"')
    if n % 2 != 0:
        hints.append(f"出错行含 {n} 个英文双引号（奇数），几乎可以确定是引号未转义或未闭合。")
    if bad_line.rstrip().endswith(",") and "Expecting property name" in msg:
        hints.append("注意：JSON 不允许尾随逗号（最后一个字段后面不能留逗号）。")
    if re.search(r"^\s*'", bad_line):
        hints.append("检测到单引号开头的键/值，JSON 要求键和字符串都用双引号。")

    hints.append("快速自检：用 `python3 -m json.tool <file>` 也能定位到同一行。")
    return hints


def die(message):
    print(message)
    sys.exit(1)


# ---------------------------------------------------------------- 结构校验
def edge_nodes(edge):
    """从一条 Mermaid 边里取出两端节点 id。

    需先剥掉两种标签写法，否则标签文字会被误当成节点：
        A -->|文字| B       A -- 文字 --> B
    """
    s = str(edge or "")
    s = re.sub(r"--\s+[^-|]+?\s+-->", "-->", s)   # A -- 文字 --> B
    s = re.sub(r"\|[^|]*\|", " ", s)              # A -->|文字| B
    parts = re.split(r"[-.=<>]{2,}", s)
    return [p.strip() for p in parts if p.strip()]


def validate_structure(data):
    """非致命检查：只警告，不阻断。"""
    warnings = []

    for key in ["paper"] + [s.lower() for s in SHEETS]:
        if key not in data:
            warnings.append(f"缺少顶层键 `{key}`（将输出空表）")

    for sheet in SHEETS:
        rows = data.get(sheet.lower(), [])
        if not isinstance(rows, list):
            warnings.append(f"`{sheet.lower()}` 不是数组")
            continue
        for i, row in enumerate(rows, start=1):
            if not isinstance(row, dict):
                warnings.append(f"{sheet} 第 {i} 行不是对象")
                continue
            unknown = [k for k in row if k not in SCHEMA[sheet]]
            if unknown:
                warnings.append(f"{sheet} 第 {i} 行有未定义字段 {unknown}（会追加到末尾）")

    # 变量名一致性（跨研究合并靠它连链）
    names = Counter()
    for row in data.get("variables", []):
        if isinstance(row, dict):
            names[row.get("variable_name_en", "")] += 1
    lowered = {}
    for name in names:
        lowered.setdefault(name.strip().lower(), []).append(name)
    for variants in lowered.values():
        if len(set(variants)) > 1:
            warnings.append(f"变量名大小写/空格不一致，跨研究合并会断链: {sorted(set(variants))}")

    # 题项数对账：标注了 n_items 却一条题项都没进 Items 表，必须显式交代
    items_by_var = Counter()
    for row in data.get("items", []):
        if isinstance(row, dict):
            items_by_var[row.get("variable_name_en", "")] += 1
    for i, row in enumerate(data.get("variables", []), start=1):
        if not isinstance(row, dict):
            continue
        name = row.get("variable_name_en", "")
        try:
            n_int = int(str(row.get("n_items")).strip())
        except (TypeError, ValueError):
            continue
        if n_int <= 0:
            continue
        table_rows = items_by_var.get(name, 0)
        given_raw = row.get("items_given")
        given_set = False
        try:
            given_int = int(str(given_raw).strip())
            given_set = True
        except (TypeError, ValueError):
            given_int = table_rows
        if given_int > n_int:
            warnings.append(
                f"Variables 第 {i} 行 {name} 的 items_given={given_int} 大于 n_items={n_int}"
            )
        elif table_rows == 0 and not given_set:
            # 没填 items_given 又一条题项都没有：无法判断是「原文未提供」还是「漏了」
            warnings.append(
                f"Variables 第 {i} 行 {name} 标注 {n_int} 题，但 Items 表没有它的题项；"
                "若原文确实只给样例题项或未给题项，请在 items_given 写实际条数"
                "（可以写 0），并在 missing_fields 记 items"
            )
        elif table_rows == 0 and given_int > 0:
            warnings.append(
                f"Variables 第 {i} 行 {name} 的 items_given={given_int}，"
                "但 Items 表里没有对应行——题项要逐条写进 Items 表"
            )

    # 信度数值：只该填 0–1 的数字
    for i, row in enumerate(data.get("variables", []), start=1):
        if not isinstance(row, dict):
            continue
        for field in ("cr", "alpha", "ave", "avg_loading"):
            raw = row.get(field)
            if raw in (None, "", "NA"):
                continue
            try:
                x = float(str(raw).strip())
            except ValueError:
                warnings.append(
                    f"Variables 第 {i} 行 {field} = `{raw}` 不是数字"
                    "（信度只填数值；文字说明写 notes）"
                )
                continue
            if not 0 <= x <= 1:
                warnings.append(f"Variables 第 {i} 行 {field} = {x} 超出 0–1 范围")

    # 计算测量：方法类型与模型文档要配套
    for i, row in enumerate(data.get("studies", []), start=1):
        if not isinstance(row, dict):
            continue
        if row.get("method_type") == "Computational Measurement":
            if row.get("model_spec") in (None, "", "NA"):
                warnings.append(
                    f"Studies 第 {i} 行标为 Computational Measurement，但缺 model_spec"
                    "（架构 / 基座 / 关键超参）"
                )
    # 分析模型：假设与命题必须成对记录
    for i, row in enumerate(data.get("studies", []), start=1):
        if not isinstance(row, dict):
            continue
        if row.get("method_type") == "Analytical Model":
            for field, need in (("key_assumptions", "关键假设"),
                                ("model_results", "命题/结果")):
                if row.get(field) in (None, "", "NA"):
                    warnings.append(
                        f"Studies 第 {i} 行标为 Analytical Model，但缺 {field}（{need}）——"
                        "命题脱离假设就没有意义，两者必须成对记录"
                    )

    # 推导类条目：supported 应写 Derived，不要写成 Supported/Observed
    for i, row in enumerate(data.get("hypotheses", []), start=1):
        if not isinstance(row, dict):
            continue
        if row.get("test_type") == "Derivation":
            sup = row.get("supported")
            if sup in ("Supported", "Partially supported", "Observed",
                       "Partially observed", "Not observed"):
                warnings.append(
                    f"Hypotheses 第 {i} 行 test_type = Derivation，但 supported = `{sup}`；"
                    "推导得到的命题应写 `Derived`（被证明 ≠ 被观察到），"
                    "实证是否一致写进 note"
                )

    # n 的单位：多研究论文里 n 可能是图像数、判断条数、品牌数、受访者数，必须逐行标明
    study_rows = [s for s in data.get("studies", []) if isinstance(s, dict)]
    if len(study_rows) >= 2:
        for i, row in enumerate(study_rows, start=1):
            if row.get("n") not in (None, "", "NA") and row.get("n_unit") in (None, "", "NA"):
                warnings.append(
                    f"Studies 第 {i} 行有 n 但缺 n_unit——多研究论文里 n 可能是"
                    "图像数、判断条数、品牌数或受访者数，必须标明单位"
                )
    has_model_spec = any(
        isinstance(s, dict) and s.get("model_spec") not in (None, "", "NA")
        for s in data.get("studies", []))
    if not has_model_spec:
        for i, row in enumerate(data.get("variables", []), start=1):
            if isinstance(row, dict) and row.get("model_performance") not in (None, "", "NA"):
                warnings.append(
                    f"Variables 第 {i} 行填了 model_performance，"
                    "但没有任何研究写 model_spec——模型细节无从追溯"
                )
                break

    # 主题层级断链：parent_theme 必须能在变量表里找到同名行
    known = {row.get("variable_name_en", "") for row in data.get("variables", [])
             if isinstance(row, dict)}
    for i, row in enumerate(data.get("variables", []), start=1):
        if not isinstance(row, dict):
            continue
        parent = row.get("parent_theme")
        level = row.get("theme_level")
        has_parent = bool(parent) and parent != "NA"
        uses_tree = bool(level) and level != "NA"
        if not has_parent:
            continue
        if not uses_tree:
            # 量化论文这两列应写 NA（codebook 第 83 行）；此时报「断链」是错误诊断
            warnings.append(
                f"Variables 第 {i} 行填了 parent_theme = `{parent}` 但 theme_level = NA："
                "量化论文这两列应写 NA（只有质性/元分析才用主题树）。"
                "若本文确为质性或元分析，请补 theme_level"
            )
        elif parent not in known:
            warnings.append(
                f"Variables 第 {i} 行 parent_theme = `{parent}` 在变量表里没有同名上级"
                "（主题层级树会断链，请检查名称是否大小写/空格一致）"
            )

    # 受控词表提示（按表检查，不同字段落在不同 sheet）
    for sheet, fields in SHEET_VOCAB.items():
        for i, row in enumerate(data.get(sheet, []), start=1):
            if not isinstance(row, dict):
                continue
            for field in fields:
                allowed = CONTROLLED_VOCAB[field]
                val = row.get(field)
                if not val or val == "NA":
                    continue
                # design 常被写成程序描述；仅在确实不在词表内时才这样提示（多值合法时不误报）
                if field == "design":
                    parts = [p.strip() for p in str(val).split("|")]
                    if any(p not in allowed for p in parts) and (
                        len(str(val)) > 24 or re.search(r"[\u4e00-\u9fff]", str(val))
                    ):
                        warnings.append(
                            f"Studies 第 {i} 行 design 写得像程序描述（`{str(val)[:24]}…`）；"
                            "design 只填设计类型（如 Between-subjects），程序细节写入 main_purpose / notes"
                        )
                        continue
                if val in allowed:
                    continue
                # 允许用 " | " 分隔的多值
                parts = [p.strip() for p in str(val).split("|")]
                bad = [p for p in parts if p not in allowed]
                if not bad:
                    continue
                off = [a for a in allowed if a.lower() == bad[0].lower()]
                if off:
                    tail = f"，是否想写 `{off[0]}`？"
                elif len(allowed) <= 10:
                    tail = f"（建议取值: {sorted(allowed)}）"
                else:
                    tail = (f"（词表共 {len(allowed)} 项，完整取值见 references/excel-codebook.md；"
                            "若这其实是估计方法而非设计，写入 estimator 列）")
                warnings.append(
                    f"{SHEET_LABEL.get(sheet, sheet)} 第 {i} 行 {field} = `{bad[0]}` 不在受控词表内" + tail
                )

    # Mermaid 边指向未定义节点：渲染器会凭空画出这个节点，必须查
    study_ids = {str(row.get("study_id", "")).strip()
                 for row in data.get("studies", []) if isinstance(row, dict)}

    edges_raw = data.get("mermaid_edges")
    if edges_raw is None:
        edges_raw = []
    elif not isinstance(edges_raw, list):
        warnings.append("mermaid_edges 应是数组（如 [\"S1 --> S2\"]），当前类型不对，已忽略")
        edges_raw = []
    groups_raw = data.get("mermaid_groups")
    if groups_raw is None:
        groups_raw = {}
    elif not isinstance(groups_raw, dict):
        warnings.append("mermaid_groups 应是对象（如 {\"主链\": [\"S1\", \"S2\"]}），当前类型不对，已忽略")
        groups_raw = {}

    for i, edge in enumerate(edges_raw or [], start=1):
        for node in edge_nodes(edge):
            if node and node not in study_ids:
                warnings.append(
                    f"mermaid_edges 第 {i} 条引用了未定义的节点 `{node}`"
                    f"（已定义的 study_id: {sorted(study_ids)}）。"
                    "Mermaid 会凭空画出这个节点——检查拼写，或先补上该研究"
                )
    for title, members in groups_raw.items():
        for sid in members or []:
            if str(sid).strip() not in study_ids:
                warnings.append(
                    f"mermaid_groups「{title}」里的 `{sid}` 不是任何研究的 study_id，"
                    "该成员会被静默丢弃"
                )

    # 空白未填 NA —— 但 missing_fields 例外：留空即表示「无缺失字段」
    for sheet in SHEETS:
        for i, row in enumerate(data.get(sheet.lower(), []), start=1):
            if isinstance(row, dict):
                blanks = [k for k, v in row.items()
                          if k not in OPTIONAL_BLANK
                          and (v is None or (isinstance(v, str) and not v.strip()))]
                if blanks:
                    warnings.append(f"{sheet} 第 {i} 行有空值字段 {blanks}（应写 NA）")

    return warnings


# ---------------------------------------------------------------- 写 Excel
def build_sheet(wb, title, rows):
    ws = wb.create_sheet(title=title)
    columns = list(SCHEMA[title])

    extra = []
    for row in rows:
        for key in row:
            if key not in columns and key not in extra:
                extra.append(key)
    columns.extend(extra)

    for col_idx, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, name in enumerate(columns, start=1):
            value = row.get(name, "")
            if value is None or value == "":
                # missing_fields 留空 = 无缺失字段，不写 NA（与 excel-codebook.md 一致）
                value = "" if name == "missing_fields" else "NA"
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.alignment = BODY_ALIGN

    ws.freeze_panes = "A2"
    for c_idx, name in enumerate(columns, start=1):
        longest = max(
            [len(str(name))]
            + [len(str(row.get(name, ""))) for row in rows]
            or [len(str(name))]
        )
        width = min(max(12, min(longest, 60) + 2), 60)
        ws.column_dimensions[get_column_letter(c_idx)].width = width
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        sys.exit(1)

    in_path = args[0]
    validate_only = "--validate" in args
    mermaid_only = "--mermaid" in args
    out_path = next((a for a in args[1:] if not a.startswith("--")), None)

    if not (validate_only or mermaid_only) and not out_path:
        print("缺少输出路径。用法: export_codebook.py <input.json> <output.xlsx>")
        sys.exit(1)

    data, _ = load_json_with_diagnostics(in_path)
    print(f"[1/3] JSON 解析通过: {in_path}")

    # 顶层键别名：effect_sizes 与 effectsizes 都接受
    if "effect_sizes" in data and "effectsizes" not in data:
        data["effectsizes"] = data["effect_sizes"]

    # 为每个研究补齐 mermaid_node 列（可直接复制进流程图）
    for s in data.get("studies", []):
        if isinstance(s, dict) and not s.get("mermaid_node"):
            s["mermaid_node"] = mm_node_for(s)

    warnings = validate_structure(data)
    if warnings:
        print(f"[2/3] 结构检查发现 {len(warnings)} 条提示:")
        for w in warnings[:40]:
            print(f"      ! {w}")
        if len(warnings) > 40:
            print(f"      ... 其余 {len(warnings) - 40} 条已省略")
    else:
        print("[2/3] 结构检查通过（无提示）")

    # ---- Mermaid 研究地图
    studies = data.get("studies", [])
    mermaid_src = build_mermaid(
        studies,
        edges=data.get("mermaid_edges"),
        groups=data.get("mermaid_groups"),
    )

    if mermaid_only:
        print("[3/3] Mermaid 研究地图源码（可直接粘进报告 1.2 节）:\n")
        print(mermaid_src)
        return

    if validate_only:
        print("[3/3] 仅校验模式，未生成文件。")
        return

    wb = Workbook()
    wb.remove(wb.active)

    info = data.get("paper", {})
    ws_info = wb.create_sheet(title="Paper")
    ws_info.append(["Field", "Value"])
    for cell in ws_info[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    labels = {
        "title": "文献标题",
        "authors": "作者",
        "year": "年份",
        "journal": "期刊",
        "doi": "DOI",
        "extracted_on": "提取日期",
        "note": "备注",
    }
    for key, label in labels.items():
        ws_info.append([label, info.get(key, "NA")])
    ws_info.column_dimensions["A"].width = 16
    ws_info.column_dimensions["B"].width = 80
    for row in ws_info.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = BODY_ALIGN

    for sheet in SHEETS:
        build_sheet(wb, sheet, data.get(sheet.lower(), []))

    wb.save(out_path)
    counts = {s: len(data.get(s.lower(), [])) for s in SHEETS}
    print(f"[3/3] 已生成: {out_path}")
    print("      " + " | ".join(f"{k}: {v} 行" for k, v in counts.items()))

    if mermaid_src:
        mmd_path = out_path.rsplit(".", 1)[0] + ".mmd"
        with open(mmd_path, "w", encoding="utf-8") as f:
            f.write(mermaid_src)
        print(f"      研究地图源码: {mmd_path}")
        print("\n--- Mermaid 研究地图（可直接粘进报告 1.2 节）---\n")
        print(mermaid_src)


if __name__ == "__main__":
    main()
