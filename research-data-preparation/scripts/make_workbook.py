#!/usr/bin/env python3
"""Create the data workbook (Excel) for a research data plan.

  python make_workbook.py --out data_workbook.xlsx                 # blank template with example rows
  python make_workbook.py --out data_workbook.xlsx --from-spec research_data_spec.json   # filled from a spec

Sheets: 说明 | Construct Map | Variable Dictionary | Sample & Rules | Data Acquisition |
        Analysis Plan | Ethics & DMP | Decision & Deviation Log | Handoff & QA
Dropdowns use the SAME English codes as research_data_spec.schema.json, so
workbook_to_spec.py can turn the workbook back into a spec without guessing.
Rows whose cells contain "（示例）" are examples and are ignored by workbook_to_spec.py.
Multi-item cells use one item per line (Alt+Enter inside a cell).
"""
import argparse, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
HDR_FILL = PatternFill("solid", start_color="1F3864")
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
EX_FONT = Font(name=FONT, size=10, italic=True, color="7F7F7F")
TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
NOTE = Font(name=FONT, size=9, italic=True, color="595959")
KEY_FILL = PatternFill("solid", start_color="F2F2F2")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

ENUMS = {
    "decision_status": ["pre_locked", "locked_after_exposure", "data_driven", "open"],
    "evidence_status": ["researcher_stated", "source_documented", "profile_verified", "needs_verification"],
    "role": ["outcome", "predictor", "moderator", "mediator", "covariate", "id", "time", "other"],
    "type": ["continuous", "binary", "categorical", "ordinal", "text", "date", "id"],
    "bool": ["TRUE", "FALSE"],
    "exposure": ["E0", "E1", "E2", "E3", "E4"],
    "design_purpose": ["confirmatory", "exploratory", "prediction", "measurement_development", "descriptive", "theory_building", "mixed"],
    "data_path": ["primary_quant", "primary_qual", "secondary_structured", "secondary_digital_trace", "mixed"],
    "readiness": ["ready", "ready_if_conditions_met", "blocked", "insufficient_evidence"],
    "ethics": ["not_required", "pending", "approved", "unknown"],
    "personal": ["none", "pseudonymised", "identifiable", "unknown"],
    "strategy": ["none", "blinded_summary", "holdout_sample", "disclosure_only"],
}

ETHICS_TOPICS = [
    ("伦理与法律", "是否需要伦理审批、知情同意、豁免审查？涉及哪些适用规则？（本表只记录，结论需人工核验）"),
    ("数据最小化", "是否确实需要收集每个字段？哪些可以不收集、删除或分开保存？"),
    ("隐私与安全", "如何匿名化或去标识化？数据如何加密？访问权限如何分配？"),
    ("保存与删除", "原始数据、处理后数据、身份对应表分别保存多久？由谁管理？"),
    ("数据使用权", "数据是否允许研究使用、保存、分享或用于后续发表？（二手数据看使用条款）"),
    ("数据共享", "能否开放原始数据、衍生数据、代码或仅汇总统计？是否需要数据使用协议？"),
    ("跨境问题", "是否涉及跨境存储、访问、传输？是否需要专项审查？（需人工核验）"),
    ("责任分工", "谁负责获取、处理、审查、授权、备份与最终归档？"),
]


def header(ws, row, cols, widths=None):
    for j, h in enumerate(cols, 1):
        c = ws.cell(row=row, column=j, value=h)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w


def put_rows(ws, start, rows, example=False):
    for i, r in enumerate(rows):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=start + i, column=j, value=v)
            c.font = EX_FONT if example else BODY
            c.border, c.alignment = BORDER, WRAP


def dropdown(ws, col_letter, first, last, key):
    dv = DataValidation(type="list", formula1='"' + ",".join(ENUMS[key]) + '"', allow_blank=True)
    dv.error, dv.errorTitle = "请从下拉列表中选择", "输入无效"
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}{first}:{col_letter}{last}")


def j(xs):
    return ", ".join(xs) if isinstance(xs, list) else (xs or "")


def build(out, spec=None):
    wb = Workbook()
    ex = spec is None
    S = spec or {}
    study, expo = S.get("study", {}), S.get("data_exposure", {})
    gov, ready = S.get("governance", {}), S.get("readiness", {})
    hold, sr = S.get("holdout_or_blinding", {}), S.get("sample_rules", {})

    # ---------- 说明
    ws = wb.active
    ws.title = "说明"
    ws["A1"] = "研究数据工作簿"
    ws["A1"].font = TITLE
    ws["A2"] = ("填写规则：下拉列表使用英文代码，与 research_data_spec.json 一致。灰色斜体且含“（示例）”的行是示例，"
                "填写前可保留或删除，转换为 JSON 时会被忽略。decision_status 说明“何时/如何定下来”，"
                "evidence_status 说明“信息核实到什么程度”，二者都不是对数据质量的评分。")
    ws["A2"].font, ws["A2"].alignment = NOTE, WRAP
    ws.merge_cells("A2:C2")
    ws.row_dimensions[2].height = 48
    header(ws, 4, ["项目", "内容", "说明"], [30, 70, 55])
    kv = [
        ("title", "研究题目", study.get("title", ""), ""),
        ("research_question", "研究问题", study.get("research_question", ""), ""),
        ("target_population", "目标人群", study.get("target_population", ""), "理论上要研究谁；与“实际样本”区分"),
        ("unit_of_analysis", "分析单位", study.get("unit_of_analysis", ""), "观测单位与分析单位是否一致？"),
        ("design_purpose", "研究目的", study.get("design_purpose", ""), "决定哪些决策必须事前锁定"),
        ("data_path", "数据路径", study.get("data_path", ""), ""),
        ("estimand", "目标估计量（大白话）", study.get("estimand", ""), "比较谁、比较什么、在什么时间范围；确认性解释/因果研究需要"),
        ("exposure_level", "数据接触等级", expo.get("level", ""), "E0未接触 E1仅元数据 E2结构/有限样本 E3分布或描述 E4假设相关分析"),
        ("exposure_date", "接触等级日期", expo.get("date", ""), "YYYY-MM-DD"),
        ("exposure_description", "接触情况说明", expo.get("description", ""), "看过什么、没看过什么"),
        ("exposure_log_ref", "接触日志文件", expo.get("log_ref", ""), "profile_dataset.py 生成的 exposure_log.jsonl 路径"),
        ("readiness_status", "准备状态", ready.get("status", ""), "是研究准备与数据获取的状态，不是数据质量评级"),
        ("readiness_conditions", "满足条件（每行一条）", "\n".join(ready.get("conditions", [])), "单元格内换行：Alt+Enter"),
        ("holdout_strategy", "盲化/留出策略", hold.get("strategy", ""), "none / blinded_summary / holdout_sample / disclosure_only"),
        ("holdout_details", "策略细节", hold.get("details", ""), ""),
        ("ethics_approval_status", "伦理审批状态", gov.get("ethics_approval_status", ""), ""),
        ("personal_data", "个人信息", gov.get("personal_data", ""), ""),
        ("human_verification", "需人工核验事项（每行一条）", "\n".join(gov.get("items_needing_human_verification", [])), "法律、伦理、数据使用条款等；单元格内换行：Alt+Enter"),
        ("storage_and_retention", "存储与保留", gov.get("storage_and_retention", ""), ""),
    ]
    rowmap = {}
    for i, (key, label, val, note) in enumerate(kv, 5):
        rowmap[key] = i
        for jx, v in enumerate((label, val, note), 1):
            c = ws.cell(row=i, column=jx, value=v)
            c.font, c.border, c.alignment = (BOLD if jx == 1 else BODY), BORDER, WRAP
        ws.cell(row=i, column=1).fill = KEY_FILL
    for key, enum in (("design_purpose", "design_purpose"), ("data_path", "data_path"), ("exposure_level", "exposure"),
                      ("readiness_status", "readiness"), ("holdout_strategy", "strategy"),
                      ("ethics_approval_status", "ethics"), ("personal_data", "personal")):
        dropdown(ws, "B", rowmap[key], rowmap[key], enum)
    r0 = 5 + len(kv) + 1
    ws.cell(row=r0, column=1, value="决策状态统计（公式，自动计算）").font = BOLD
    stat_rows = [
        ("变量：事前锁定", "=COUNTIF('Variable Dictionary'!I:I,\"pre_locked\")"),
        ("变量：接触后锁定", "=COUNTIF('Variable Dictionary'!I:I,\"locked_after_exposure\")"),
        ("变量：数据驱动", "=COUNTIF('Variable Dictionary'!I:I,\"data_driven\")"),
        ("变量：待定", "=COUNTIF('Variable Dictionary'!I:I,\"open\")"),
        ("变量：证据待核实", "=COUNTIF('Variable Dictionary'!J:J,\"needs_verification\")"),
        ("决策日志：条数", "=COUNTA('Decision & Deviation Log'!A5:A200)-COUNTIF('Decision & Deviation Log'!A5:A200,\"*示例*\")"),
    ]
    for i, (lab, f) in enumerate(stat_rows, r0 + 1):
        ws.cell(row=i, column=1, value=lab).font = BODY
        c = ws.cell(row=i, column=2, value=f)
        c.font, c.alignment = BODY, Alignment(horizontal="left")
        ws.cell(row=i, column=1).border = ws.cell(row=i, column=2).border = BORDER
    ws.cell(row=r0 + len(stat_rows) + 1, column=1,
            value="统计只计入已填写的行；示例行的状态写作“（示例）”，不会被计数。").font = NOTE

    # ---------- Construct Map
    ws = wb.create_sheet("Construct Map")
    ws["A1"] = "构念—测量—字段映射"
    ws["A1"].font = TITLE
    ws["A2"] = "每个核心构念一行：理论定义 → 操作化指标（变量ID，逗号分隔）。指标名称不等于构念本身。"
    ws["A2"].font = NOTE
    header(ws, 4, ["id", "构念名称", "理论/操作定义", "指标（变量ID，逗号分隔）", "evidence_status"], [10, 24, 50, 32, 20])
    if ex:
        put_rows(ws, 5, [["（示例）C1", "品牌忠诚", "消费者对品牌的重复选择倾向", "V_repeat, V_share", "（示例）"]], True)
        last = 40
    else:
        rows = [[c["id"], c["name"], c["definition"], j(c["indicators"]), c.get("evidence_status", "")] for c in S.get("constructs", [])]
        put_rows(ws, 5, rows)
        last = max(40, 5 + len(rows) + 20)
    dropdown(ws, "E", 5, last, "evidence_status")

    # ---------- Variable Dictionary
    ws = wb.create_sheet("Variable Dictionary")
    ws["A1"] = "变量字典"
    ws["A1"].font = TITLE
    ws["A2"] = "role/type/decision_status/evidence_status 请从下拉列表选择；source_field 与 derivation 写清楚原始字段与计算方式。"
    ws["A2"].font = NOTE
    header(ws, 4, ["id", "label", "role", "type", "unit", "source_field", "derivation", "missing_codes",
                   "decision_status", "evidence_status"], [12, 28, 12, 12, 10, 18, 42, 14, 22, 20])
    if ex:
        put_rows(ws, 5, [["（示例）V_repeat", "复购次数（观察期）", "outcome", "continuous", "count", "order_id",
                          "观察期内同一customer_id的订单数", "无", "（示例）", "（示例）"]], True)
        last = 200
    else:
        rows = [[v["id"], v["label"], v["role"], v["type"], v.get("unit", ""), v.get("source_field", ""),
                 v.get("derivation", ""), v.get("missing_codes", ""), v["decision_status"], v["evidence_status"]]
                for v in S.get("variables", [])]
        put_rows(ws, 5, rows)
        last = max(200, 5 + len(rows) + 50)
    for col, key in (("C", "role"), ("D", "type"), ("I", "decision_status"), ("J", "evidence_status")):
        dropdown(ws, col, 5, last, key)

    # ---------- Sample & Rules
    ws = wb.create_sheet("Sample & Rules")
    ws["A1"] = "抽样框架与样本规则"
    ws["A1"].font = TITLE
    header(ws, 3, ["项目", "内容", "说明"], [28, 70, 45])
    tw = sr.get("time_window", {})
    skv = [
        ("frame", "抽样框架 / 数据实际覆盖谁", sr.get("frame", ""), "区分：理论目标人群 vs 可接触总体 vs 实际样本"),
        ("inclusion", "纳入标准（每行一条）", "\n".join(sr.get("inclusion", [])), "单元格内换行：Alt+Enter"),
        ("window_start", "时间窗口：开始", tw.get("start", ""), "二手数据：先于结果确定，并写明理由"),
        ("window_end", "时间窗口：结束", tw.get("end", ""), ""),
        ("window_rationale", "时间窗口理由", tw.get("rationale", ""), "是否避开结构性断裂？"),
        ("stopping_rule", "停止规则（一手数据）", sr.get("stopping_rule", ""), "何时停止收集；不得看了结果再定"),
        ("decision_status", "整体决策状态", sr.get("decision_status", ""), ""),
    ]
    smap = {}
    for i, (key, label, val, note) in enumerate(skv, 4):
        smap[key] = i
        for jx, v in enumerate((label, val, note), 1):
            c = ws.cell(row=i, column=jx, value=v)
            c.font, c.border, c.alignment = (BOLD if jx == 1 else BODY), BORDER, WRAP
        ws.cell(row=i, column=1).fill = KEY_FILL
    dropdown(ws, "B", smap["decision_status"], smap["decision_status"], "decision_status")
    ws.cell(row=12, column=1, value="排除规则").font = BOLD
    header(ws, 13, ["规则", "decided_before_outcome", "说明"])
    if ex:
        put_rows(ws, 14, [["（示例）未通过注意力检查", "（示例）", "TRUE 表示在看到结果数据之前就已确定"]], True)
        lastx = 40
    else:
        rows = [[e["rule"], str(e["decided_before_outcome"]).upper(), ""] for e in sr.get("exclusion", [])]
        put_rows(ws, 14, rows)
        lastx = max(40, 14 + len(rows) + 15)
    dropdown(ws, "B", 14, lastx, "bool")

    # ---------- Data Acquisition
    ws = wb.create_sheet("Data Acquisition")
    ws["A1"] = "数据来源与获取记录"
    ws["A1"].font = TITLE
    ws["A2"] = ("区分三件事：官方是否有该数据产品 / 当前账户是否实际有访问权限 / 该数据是否支持本研究。"
                "每项都要有核验日期，不确定的写“未核实”。")
    ws["A2"].font = NOTE
    header(ws, 4, ["来源名称", "官方入口", "数据内容（变量/时间/地域/单位）", "获取方式", "使用条件（许可/费用/分享限制）",
                   "产品存在已核验？", "账户有访问权限？", "支持本研究已核验？", "已知限制", "版本/提取日期", "核验日期"],
           [22, 28, 34, 16, 30, 14, 14, 16, 30, 16, 14])
    if ex:
        put_rows(ws, 5, [["（示例）某公开数据集", "https://…", "交易记录；2018-2020；单位：订单", "公开下载",
                          "仅限研究使用，禁止再分发", "是（2026-10-03）", "是", "未核实", "2020年记录量异常", "v1 / 2026-10-03", "2026-10-03"]], True)

    # ---------- Analysis Plan
    ws = wb.create_sheet("Analysis Plan")
    ws["A1"] = "分析计划"
    ws["A1"].font = TITLE
    ws["A2"] = ("primary=TRUE 的是确认性主要分析。缺失、离群值、多重检验、稳健性的预案也写在这里；"
                "实际执行交给 quant-data-analysis 等分析skill。")
    ws["A2"].font = NOTE
    header(ws, 4, ["id", "假设/问题", "模型与检验", "primary", "decision_status", "备注"], [10, 44, 48, 10, 22, 40])
    if ex:
        put_rows(ws, 5, [["（示例）A1", "H1：品类广度正向预测新增品类数", "Poisson回归，稳健标准误", "（示例）", "（示例）", "同时报告机械上限检查"]], True)
        last = 60
    else:
        rows = [[a["id"], a["hypothesis"], a["model"], str(a["primary"]).upper(), a["decision_status"], a.get("notes", "")]
                for a in S.get("analysis_plan", [])]
        put_rows(ws, 5, rows)
        last = max(60, 5 + len(rows) + 20)
    dropdown(ws, "D", 5, last, "bool")
    dropdown(ws, "E", 5, last, "decision_status")

    # ---------- Ethics & DMP
    ws = wb.create_sheet("Ethics & DMP")
    ws["A1"] = "伦理与数据管理计划（DMP）"
    ws["A1"].font = TITLE
    ws["A2"] = "本表只记录问题与处理情况，不给出法律或伦理结论；涉及个人信息、跨境、平台规则的事项需人工或专业核验。"
    ws["A2"].font = NOTE
    header(ws, 4, ["主题", "需要回答的问题", "当前答案/处理方式", "是否已核验（人工）", "负责人"], [18, 60, 44, 18, 14])
    put_rows(ws, 5, [[t, q, "", "", ""] for t, q in ETHICS_TOPICS])

    # ---------- Decision & Deviation Log
    ws = wb.create_sheet("Decision & Deviation Log")
    ws["A1"] = "决策与偏离日志"
    ws["A1"].font = TITLE
    ws["A2"] = "每个重要决策一行；之后的调整新增一行，不覆盖旧行。写清当时的接触等级。"
    ws["A2"].font = NOTE
    header(ws, 4, ["id", "date", "decision", "basis", "status", "previous", "reason_for_change", "exposure_at_time"],
           [10, 12, 40, 36, 22, 28, 32, 16])
    if ex:
        put_rows(ws, 5, [["（示例）D1", "2026-10-03", "分析窗口截止到2019年底", "结构级剖面显示2020年记录量异常",
                          "（示例）", "", "", "E2"]], True)
        last = 200
    else:
        rows = [[d["id"], d["date"], d["decision"], d.get("basis", ""), d["status"], d.get("previous", ""),
                 d.get("reason_for_change", ""), d.get("exposure_at_time", "")] for d in S.get("decision_log", [])]
        put_rows(ws, 5, rows)
        last = max(200, 5 + len(rows) + 50)
    dropdown(ws, "E", 5, last, "decision_status")
    dropdown(ws, "H", 5, last, "exposure")

    # ---------- Handoff & QA
    ws = wb.create_sheet("Handoff & QA")
    ws["A1"] = "交接与验收"
    ws["A1"].font = TITLE
    header(ws, 3, ["接收方（skill或agent）", "交付内容（逗号分隔）", "状态", "验收标准/未解决风险"], [34, 50, 22, 48])
    if ex:
        put_rows(ws, 4, [["（示例）quant-data-analysis", "research_data_spec.json, data_workbook.xlsx", "（示例）", "接触等级与排除规则已写明"]], True)
    else:
        rows = [[d["skill_or_agent"], j(d["inputs"]), d["status"], ""] for d in S.get("handoff", {}).get("downstream", [])]
        put_rows(ws, 4, rows)

    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
    wb.save(out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--from-spec", default=None)
    a = ap.parse_args()
    spec = json.load(open(a.from_spec, encoding="utf-8")) if a.from_spec else None
    build(a.out, spec)
    print(f"Saved {a.out}" + (f" (filled from {a.from_spec})" if spec else " (blank template)"))
