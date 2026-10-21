#!/usr/bin/env python3
"""Convert the data workbook (made by make_workbook.py) into research_data_spec.json.

  python workbook_to_spec.py data_workbook.xlsx --out research_data_spec.json

The workbook is the single place humans edit; the JSON is generated from it so the two
cannot drift apart. Rows containing "（示例）" are skipped. Afterwards run:
  python validate_spec.py research_data_spec.json
"""
import argparse, datetime as dt, json
from openpyxl import load_workbook


def s(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def is_example(row):
    return any("（示例）" in s(c) for c in row)


def table(ws, first_row, ncols, key_col=1):
    out = []
    for r in ws.iter_rows(min_row=first_row, max_col=ncols, values_only=True):
        if not s(r[key_col - 1]) or is_example(r):
            continue
        out.append([s(c) for c in r])
    return out


def split(v, sep):
    return [x.strip() for x in v.split(sep) if x.strip()]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    wb = load_workbook(a.workbook, data_only=True)

    info = {}
    for r in wb["说明"].iter_rows(min_row=5, max_col=2, values_only=True):
        if s(r[0]):
            info[s(r[0])] = s(r[1])
    g = lambda label: info.get(label, "")

    study = {"title": g("研究题目"), "research_question": g("研究问题"), "target_population": g("目标人群"),
             "unit_of_analysis": g("分析单位"), "design_purpose": g("研究目的"), "data_path": g("数据路径")}
    if g("目标估计量（大白话）"):
        study["estimand"] = g("目标估计量（大白话）")
    exposure = {"level": g("数据接触等级"), "date": g("接触等级日期"), "description": g("接触情况说明")}
    if g("接触日志文件"):
        exposure["log_ref"] = g("接触日志文件")
    readiness = {"status": g("准备状态"), "conditions": split(g("满足条件（每行一条）"), "\n")}
    holdout = {"strategy": g("盲化/留出策略") or "none"}
    if g("策略细节"):
        holdout["details"] = g("策略细节")
    gov = {"ethics_approval_status": g("伦理审批状态"), "personal_data": g("个人信息"),
           "items_needing_human_verification": split(g("需人工核验事项（每行一条）"), "\n")}
    if g("存储与保留"):
        gov["storage_and_retention"] = g("存储与保留")

    constructs = []
    for r in table(wb["Construct Map"], 5, 5):
        c = {"id": r[0], "name": r[1], "definition": r[2], "indicators": split(r[3], ",")}
        if r[4]:
            c["evidence_status"] = r[4]
        constructs.append(c)

    variables = []
    for r in table(wb["Variable Dictionary"], 5, 10):
        v = {"id": r[0], "label": r[1], "role": r[2], "type": r[3], "decision_status": r[8], "evidence_status": r[9]}
        for key, val in (("unit", r[4]), ("source_field", r[5]), ("derivation", r[6]), ("missing_codes", r[7])):
            if val:
                v[key] = val
        variables.append(v)

    ws = wb["Sample & Rules"]
    sk = {}
    for r in ws.iter_rows(min_row=4, max_row=10, max_col=2, values_only=True):
        sk[s(r[0])] = s(r[1])
    rules = {"frame": sk.get("抽样框架 / 数据实际覆盖谁", ""), "inclusion": split(sk.get("纳入标准（每行一条）", ""), "\n"),
             "exclusion": [{"rule": r[0], "decided_before_outcome": r[1].upper() == "TRUE"}
                           for r in table(ws, 14, 2)],
             "decision_status": sk.get("整体决策状态", "")}
    tw = {k: sk.get(lbl, "") for k, lbl in (("start", "时间窗口：开始"), ("end", "时间窗口：结束"), ("rationale", "时间窗口理由"))}
    tw = {k: v for k, v in tw.items() if v}
    if tw:
        rules["time_window"] = tw
    if sk.get("停止规则（一手数据）"):
        rules["stopping_rule"] = sk["停止规则（一手数据）"]

    plan = []
    for r in table(wb["Analysis Plan"], 5, 6):
        p = {"id": r[0], "hypothesis": r[1], "model": r[2], "primary": r[3].upper() == "TRUE", "decision_status": r[4]}
        if r[5]:
            p["notes"] = r[5]
        plan.append(p)

    log = []
    for r in table(wb["Decision & Deviation Log"], 5, 8):
        d = {"id": r[0], "date": r[1], "decision": r[2], "status": r[4]}
        for key, val in (("basis", r[3]), ("previous", r[5]), ("reason_for_change", r[6]), ("exposure_at_time", r[7])):
            if val:
                d[key] = val
        log.append(d)

    handoff = {"downstream": [{"skill_or_agent": r[0], "inputs": split(r[1], ","), "status": r[2]}
                              for r in table(wb["Handoff & QA"], 4, 4)]}

    spec = {"spec_version": "1.0", "study": study, "data_exposure": exposure,
            "variables": variables, "sample_rules": rules, "analysis_plan": plan, "holdout_or_blinding": holdout,
            "governance": gov, "readiness": readiness, "decision_log": log, "handoff": handoff}
    if constructs:
        spec["constructs"] = constructs
    json.dump(spec, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    print(f"Wrote {a.out}. Next: python validate_spec.py {a.out}")


if __name__ == "__main__":
    main()
